from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import CaseConfig, resolve_path

REQUIRED_COLUMNS = [
    "Story",
    "Label",
    "UniqueName",
    "DesignSect",
    "Station",
    "AsTop",
    "AsBot",
    "VRebar",
    "TLngRebar",
    "TTrnRebar",
]


@dataclass(frozen=True)
class EtabsStationRow:
    story: str
    label: str
    unique_name: str
    design_sect: str
    station: float
    as_top: float
    as_bot: float
    v_rebar_req: float
    t_lng_req: float
    t_trn_req: float
    source_row: int | None = None


@dataclass(frozen=True)
class EtabsFrameData:
    unique_name: str
    stations: tuple[EtabsStationRow, ...]


@dataclass(frozen=True)
class EtabsSourceData:
    source_name: str
    by_unique_name: dict[str, EtabsFrameData]
    row_count: int
    skipped_rows: int
    warnings: tuple[str, ...]


def load_etabs_sources(config: CaseConfig, case_file_dir: Path) -> dict[str, EtabsSourceData]:
    seismic_path = resolve_path(config.inputs.seismic_excel, case_file_dir)
    gravity_path = resolve_path(config.inputs.gravity_excel, case_file_dir)
    unit_factor = 100.0 if config.units.rebar_per_length == "cm2/m" else 1.0

    seismic = read_etabs_excel(
        source_name="seismic",
        excel_path=seismic_path,
        sheet_name=config.inputs.sheet_name,
        rebar_per_length_factor=unit_factor,
    )
    gravity = read_etabs_excel(
        source_name="gravity",
        excel_path=gravity_path,
        sheet_name=config.inputs.sheet_name,
        rebar_per_length_factor=unit_factor,
    )
    return {"seismic": seismic, "gravity": gravity}


def read_etabs_excel(
    *,
    source_name: str,
    excel_path: Path,
    sheet_name: str,
    rebar_per_length_factor: float,
) -> EtabsSourceData:
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found for {source_name}: {excel_path}")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {excel_path}")
    ws = wb[sheet_name]

    header_row_idx, header_map = find_header(ws.iter_rows(min_row=1, max_row=30, values_only=True))
    missing = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing:
        raise ValueError(
            f"Missing required columns in {excel_path.name}: {', '.join(missing)}"
        )

    raw_by_unique: dict[str, list[EtabsStationRow]] = {}
    seen_stations: dict[str, set[float]] = {}
    duplicate_counts: dict[str, int] = {}
    row_count = 0
    skipped_rows = 0
    warnings: list[str] = []
    for row_index, row in enumerate(ws.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
        if row is None:
            continue
        unique_name = as_text(row[header_map["UniqueName"]])
        station_value = row[header_map["Station"]]
        if not unique_name or station_value in (None, ""):
            skipped_rows += 1
            continue
        try:
            station = float(station_value)
        except (TypeError, ValueError):
            skipped_rows += 1
            continue

        row_count += 1
        record = EtabsStationRow(
            story=as_text(row[header_map["Story"]]),
            label=as_text(row[header_map["Label"]]),
            unique_name=unique_name,
            design_sect=as_text(row[header_map["DesignSect"]]),
            station=station,
            as_top=as_float(row[header_map["AsTop"]]),
            as_bot=as_float(row[header_map["AsBot"]]),
            v_rebar_req=as_float(row[header_map["VRebar"]]) * rebar_per_length_factor,
            t_lng_req=as_float(row[header_map["TLngRebar"]]),
            t_trn_req=as_float(row[header_map["TTrnRebar"]]) * rebar_per_length_factor,
            source_row=row_index,
        )

        known = seen_stations.setdefault(unique_name, set())
        if station in known:
            duplicate_counts[unique_name] = duplicate_counts.get(unique_name, 0) + 1
        known.add(station)
        raw_by_unique.setdefault(unique_name, []).append(record)

    by_unique_name = {
        unique_name: EtabsFrameData(
            unique_name=unique_name,
            stations=tuple(sorted(rows, key=lambda item: (item.station, item.source_row or 0))),
        )
        for unique_name, rows in raw_by_unique.items()
    }

    if skipped_rows > 0:
        warnings.append(
            f"{source_name}: skipped {skipped_rows} rows with missing/invalid UniqueName or Station"
        )
    if duplicate_counts:
        duplicate_summary = ", ".join(
            f"{unique_name}={count}" for unique_name, count in sorted(duplicate_counts.items())
        )
        warnings.append(
            f"{source_name}: duplicate station rows preserved as independent physical scenarios "
            f"because their ETABS semantics are ambiguous ({duplicate_summary}); source_row retains provenance"
        )

    result = EtabsSourceData(
        source_name=source_name,
        by_unique_name=by_unique_name,
        row_count=row_count,
        skipped_rows=skipped_rows,
        warnings=tuple(warnings),
    )
    wb.close()
    return result


def find_header(rows: Iterable[tuple[object, ...]]) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows, start=1):
        row_map = {as_text(value): i for i, value in enumerate(row) if as_text(value)}
        if "UniqueName" in row_map and "Station" in row_map and "VRebar" in row_map:
            return idx, row_map
    raise ValueError("Unable to locate ETABS header row")


def as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def as_float(value: object) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0
