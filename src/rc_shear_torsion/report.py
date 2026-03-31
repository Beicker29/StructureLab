from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook

from .design import BeamSummary, RegionDesignResult, SpanSummary
from .results_model import to_canonical_region_result


SUPPORT_FACE_STIRRUP_OFFSET_MM = 50.0


def _region_sort_key(region_id: str) -> tuple[int, str]:
    text = str(region_id or "")
    digits = "".join(char for char in text if char.isdigit())
    return (int(digits) if digits else 10**9, text)


def _build_span_edge_regions(results: Iterable[RegionDesignResult]) -> dict[tuple[str, str], tuple[str, str]]:
    span_regions: dict[tuple[str, str], set[str]] = {}
    for result in results:
        key = (result.beam_id, result.span_id)
        span_regions.setdefault(key, set()).add(result.region_id)

    output: dict[tuple[str, str], tuple[str, str]] = {}
    for key, region_ids in span_regions.items():
        ordered = sorted(region_ids, key=_region_sort_key)
        if ordered:
            output[key] = (ordered[0], ordered[-1])
    return output


def _stirrup_count_for_reporting(
    region_length_mm: float,
    spacing_mm: int,
    *,
    is_first_region: bool,
    is_last_region: bool,
) -> int:
    if spacing_mm <= 0 or region_length_mm <= 0.0:
        return 0

    offset_start_mm = SUPPORT_FACE_STIRRUP_OFFSET_MM if is_first_region else 0.0
    offset_end_mm = SUPPORT_FACE_STIRRUP_OFFSET_MM if is_last_region else 0.0
    effective_length_mm = max(0.0, float(region_length_mm) - offset_start_mm - offset_end_mm)

    # En el vano completo se fija el primero a +50 mm (cara apoyo izq.)
    # y el ultimo a -50 mm (cara izq. apoyo der.).
    if is_first_region and is_last_region:
        return max(2, math.ceil(effective_length_mm / float(spacing_mm)) + 1)

    return max(1, math.floor(effective_length_mm / float(spacing_mm)) + 1)

DESIGN_COLUMNS = [
    "beam_id",
    "span_id",
    "region_id",
    "region_type",
    "source_control",
    "governing_station",
    "VRebar_req",
    "VRebar_req_units",
    "TTrnRebar_req",
    "TTrnRebar_req_units",
    "TLngRebar_req",
    "TLngRebar_req_units",
    "E_bar",
    "G_bar",
    "G_count",
    "spacing_mm",
    "controlling_limit",
    "Av1",
    "Av2",
    "Av_total",
    "At",
    "At_over_s",
    "Av_over_s",
    "long_bar",
    "long_count",
    "long_provided_mm2",
    "check_torsion",
    "check_shear",
    "check_longitudinal",
    "check_detailing",
    "failure_mode",
    "status",
    "message",
    "base_long_bar",
    "base_long_count",
    "extra_long_bar",
    "extra_long_count",
    "is_deep_beam",
    "longitudinal_mode",
    "longitudinal_arrangement",
]



def ensure_output_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_design_results(path: Path, region_results: Iterable[RegionDesignResult]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "design_results"
    sheet.append(DESIGN_COLUMNS)
    for result in region_results:
        canonical = to_canonical_region_result(result)
        sheet.append(
            [
                canonical.beam_id,
                canonical.span_id,
                canonical.region_id,
                canonical.region_type,
                canonical.source_control,
                canonical.governing_station,
                canonical.v_req,
                "mm2/m",
                canonical.t_req,
                "mm2/m",
                canonical.l_req,
                "mm2",
                canonical.e_bar,
                canonical.g_bar,
                canonical.g_count,
                canonical.spacing_mm,
                canonical.controlling_limit,
                canonical.av1,
                canonical.av2,
                canonical.av_total,
                canonical.at,
                canonical.at_over_s,
                canonical.av_over_s,
                canonical.long_bar,
                canonical.long_count,
                canonical.long_provided_mm2,
                canonical.checks.torsion,
                canonical.checks.shear,
                canonical.checks.longitudinal,
                canonical.checks.detailing,
                canonical.failure_mode,
                canonical.status,
                canonical.message,
                canonical.base_long_bar,
                canonical.base_long_count,
                canonical.extra_long_bar,
                canonical.extra_long_count,
                canonical.is_deep_beam,
                canonical.longitudinal_mode,
                canonical.longitudinal_arrangement_label,
            ]
        )
    workbook.save(path)


def write_summary(path: Path, span_summaries: Iterable[SpanSummary], beam_summaries: Iterable[BeamSummary]) -> None:
    workbook = Workbook()
    span_sheet = workbook.active
    span_sheet.title = "span_summary"
    span_sheet.append(
        ["beam_id", "span_id", "total_regions", "ok_regions", "fail_regions", "status", "message"]
    )
    for summary in span_summaries:
        span_sheet.append(
            [
                summary.beam_id,
                summary.span_id,
                summary.total_regions,
                summary.ok_regions,
                summary.fail_regions,
                summary.status,
                summary.message,
            ]
        )

    beam_sheet = workbook.create_sheet("beam_summary")
    beam_sheet.append(["beam_id", "total_spans", "ok_spans", "fail_spans", "status"])
    for summary in beam_summaries:
        beam_sheet.append(
            [summary.beam_id, summary.total_spans, summary.ok_spans, summary.fail_spans, summary.status]
        )

    workbook.save(path)


def write_optimized_results(path: Path, region_results: Iterable[RegionDesignResult]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "optimized_regions"
    sheet.append(
        [
            "beam_id",
            "span_id",
            "region_id",
            "method",
            "status",
            "failure_mode",
            "objective",
            "source_control",
            "governing_station",
            "E_bar",
            "G_bar",
            "G_count",
            "spacing_mm",
            "controlling_limit",
            "long_bar",
            "long_count",
            "long_provided_mm2",
            "VRebar_req",
            "TTrnRebar_req",
            "TLngRebar_req",
            "VRebar_req_units",
            "TTrnRebar_req_units",
            "TLngRebar_req_units",
            "check_torsion",
            "check_shear",
            "check_longitudinal",
            "check_detailing",
            "transverse_weight_kg_per_m",
            "longitudinal_weight_kg_per_m",
            "total_weight_kg_per_m",
            "evaluated_candidates",
            "feasible_candidates",
            "base_long_bar",
            "base_long_count",
            "extra_long_bar",
            "extra_long_count",
            "is_deep_beam",
            "longitudinal_mode",
            "longitudinal_arrangement",
        ]
    )
    for result in region_results:
        canonical = to_canonical_region_result(result)
        sheet.append(
            [
                canonical.beam_id,
                canonical.span_id,
                canonical.region_id,
                canonical.method,
                canonical.status,
                canonical.failure_mode,
                canonical.objective,
                canonical.source_control,
                canonical.governing_station,
                canonical.e_bar,
                canonical.g_bar,
                canonical.g_count,
                canonical.spacing_mm,
                canonical.controlling_limit,
                canonical.long_bar,
                canonical.long_count,
                canonical.long_provided_mm2,
                canonical.v_req,
                canonical.t_req,
                canonical.l_req,
                "mm2/m",
                "mm2/m",
                "mm2",
                canonical.checks.torsion,
                canonical.checks.shear,
                canonical.checks.longitudinal,
                canonical.checks.detailing,
                canonical.transverse_weight_kg_per_m,
                canonical.longitudinal_weight_kg_per_m,
                canonical.total_weight_kg_per_m,
                canonical.evaluated_candidates,
                canonical.feasible_candidates,
                canonical.base_long_bar,
                canonical.base_long_count,
                canonical.extra_long_bar,
                canonical.extra_long_count,
                canonical.is_deep_beam,
                canonical.longitudinal_mode,
                canonical.longitudinal_arrangement_label,
            ]
        )
    workbook.save(path)


def write_reinforcement_schedule(
    path: Path,
    region_results: Iterable[RegionDesignResult],
    region_lengths_mm: dict[tuple[str, str, str], float] | None = None,
    region_alternatives: dict[tuple[str, str, str], list[RegionDesignResult]] | None = None,
    transverse_alternatives: dict[tuple[str, str, str], list[RegionDesignResult]] | None = None,
) -> None:
    lengths_mm = region_lengths_mm or {}
    alternatives = region_alternatives or {}
    transverse_by_region = transverse_alternatives or {}
    workbook = Workbook()
    region_sheet = workbook.active
    region_sheet.title = "por_region"
    region_sheet.append(
        [
            "viga_id",
            "vano_id",
            "region_id",
            "opcion",
            "longitud_region_mm",
            "estado",
            "arreglo_transversal",
            "limite_controlante",
            "cantidad_estribos_region",
            "arreglo_longitudinal",
            "arreglo_longitudinal_base",
            "arreglo_longitudinal_adicional",
            "base_long_bar",
            "base_long_count",
            "extra_long_bar",
            "extra_long_count",
            "peso_unitario_estribo_kg",
            "peso_transversal_region_kg",
            "peso_longitudinal_region_kg",
            "peso_total_region_kg",
        ]
    )

    span_acc: dict[tuple[str, str], dict[str, float | int]] = {}
    ordered_results = sorted(region_results, key=lambda result: (result.beam_id, result.span_id, _region_sort_key(result.region_id)))
    span_edge_regions = _build_span_edge_regions(ordered_results)
    for base_result in ordered_results:
        key_region = (base_result.beam_id, base_result.span_id, base_result.region_id)
        options = alternatives.get(key_region) or [base_result]
        best_option_data: tuple[float, float, float] | None = None

        for option_idx, result in enumerate(options, start=1):
            transverse_arrangement = ""
            cantidad_estribos_region = 0
            if result.e_bar and result.spacing_mm > 0:
                if result.g_bar and result.g_count > 0:
                    transverse_arrangement = f"1E {result.e_bar} + {result.g_count}G {result.g_bar} @ {result.spacing_mm} mm"
                else:
                    transverse_arrangement = f"1E {result.e_bar} @ {result.spacing_mm} mm"

            longitudinal_arrangement = "no se requiere"
            if result.longitudinal_arrangement_label:
                longitudinal_arrangement = result.longitudinal_arrangement_label
            elif result.long_bar and result.long_count > 0:
                longitudinal_arrangement = f"{result.long_count} x {result.long_bar}"

            base_long_arrangement = "no se requiere"
            if result.base_long_bar and result.base_long_count and result.base_long_count > 0:
                base_long_arrangement = f"{result.base_long_count} x {result.base_long_bar}"
            extra_long_arrangement = "no se requiere"
            if result.extra_long_bar and result.extra_long_count and result.extra_long_count > 0:
                extra_long_arrangement = f"{result.extra_long_count} x {result.extra_long_bar}"

            estado = "cumple" if result.status == "ok" else "falla"
            longitud_region_mm = lengths_mm.get((result.beam_id, result.span_id, result.region_id), 0.0)
            span_edge = span_edge_regions.get((result.beam_id, result.span_id))
            is_first_region = bool(span_edge and result.region_id == span_edge[0])
            is_last_region = bool(span_edge and result.region_id == span_edge[1])
            if result.spacing_mm > 0 and longitud_region_mm > 0.0:
                # Estribo: conjunto formado por estribo cerrado exterior + ganchos.
                cantidad_estribos_region = _stirrup_count_for_reporting(
                    longitud_region_mm,
                    result.spacing_mm,
                    is_first_region=is_first_region,
                    is_last_region=is_last_region,
                )
            longitud_region_m = max(0.0, longitud_region_mm / 1000.0)
            peso_unitario_estribo_kg = result.stirrup_unit_weight_kg
            peso_transversal_region_kg = peso_unitario_estribo_kg * cantidad_estribos_region
            peso_longitudinal_region_kg = result.longitudinal_weight_kg_per_m * longitud_region_m
            total_weight = peso_transversal_region_kg + peso_longitudinal_region_kg
            region_sheet.append(
                [
                    result.beam_id,
                    result.span_id,
                    result.region_id,
                    option_idx,
                    longitud_region_mm,
                    estado,
                    transverse_arrangement,
                    result.controlling_limit,
                    cantidad_estribos_region,
                    longitudinal_arrangement,
                    base_long_arrangement,
                    extra_long_arrangement,
                    result.base_long_bar,
                    result.base_long_count,
                    result.extra_long_bar,
                    result.extra_long_count,
                    peso_unitario_estribo_kg,
                    peso_transversal_region_kg,
                    peso_longitudinal_region_kg,
                    total_weight,
                ]
            )
            if option_idx == 1:
                best_option_data = (
                    peso_transversal_region_kg,
                    peso_longitudinal_region_kg,
                    total_weight,
                )

        key = (base_result.beam_id, base_result.span_id)
        if key not in span_acc:
            span_acc[key] = {
                "total_regions": 0,
                "ok_regions": 0,
                "fail_regions": 0,
                "transverse_weight_sum_kg": 0.0,
                "longitudinal_weight_sum_kg": 0.0,
                "total_weight_sum_kg": 0.0,
            }
        span_row = span_acc[key]
        span_row["total_regions"] = int(span_row["total_regions"]) + 1
        if base_result.status == "ok":
            span_row["ok_regions"] = int(span_row["ok_regions"]) + 1
        else:
            span_row["fail_regions"] = int(span_row["fail_regions"]) + 1
        if best_option_data is not None:
            span_row["transverse_weight_sum_kg"] = float(span_row["transverse_weight_sum_kg"]) + best_option_data[0]
            span_row["longitudinal_weight_sum_kg"] = float(span_row["longitudinal_weight_sum_kg"]) + best_option_data[1]
            span_row["total_weight_sum_kg"] = float(span_row["total_weight_sum_kg"]) + best_option_data[2]

    span_sheet = workbook.create_sheet("por_vano")
    span_sheet.append(
        [
            "viga_id",
            "vano_id",
            "regiones_totales",
            "regiones_cumplen",
            "regiones_fallan",
            "peso_transversal_total_kg",
            "peso_longitudinal_total_kg",
            "peso_total_kg",
        ]
    )
    for beam_id, span_id in sorted(span_acc.keys()):
        span_row = span_acc[(beam_id, span_id)]
        span_sheet.append(
            [
                beam_id,
                span_id,
                int(span_row["total_regions"]),
                int(span_row["ok_regions"]),
                int(span_row["fail_regions"]),
                float(span_row["transverse_weight_sum_kg"]),
                float(span_row["longitudinal_weight_sum_kg"]),
                float(span_row["total_weight_sum_kg"]),
            ]
        )

    if transverse_by_region:
        transverse_sheet = workbook.create_sheet("transversales")
        transverse_sheet.append(
            [
                "viga_id",
                "vano_id",
                "region_id",
                "opcion",
                "estado",
                "arreglo_transversal",
                "limite_controlante",
                "espaciamiento_mm",
                "cantidad_estribos_region",
                "peso_unitario_estribo_kg",
                "peso_transversal_region_kg",
            ]
        )

        for beam_id, span_id, region_id in sorted(transverse_by_region.keys()):
            region_length_mm = lengths_mm.get((beam_id, span_id, region_id), 0.0)
            span_edge = span_edge_regions.get((beam_id, span_id))
            is_first_region = bool(span_edge and region_id == span_edge[0])
            is_last_region = bool(span_edge and region_id == span_edge[1])
            feasible_rows = [row for row in (transverse_by_region.get((beam_id, span_id, region_id)) or []) if row.status == "ok"]
            rows_sorted = sorted(
                feasible_rows,
                key=lambda row: (
                    float(row.stirrup_unit_weight_kg or 0.0)
                    * _stirrup_count_for_reporting(
                        region_length_mm,
                        row.spacing_mm,
                        is_first_region=is_first_region,
                        is_last_region=is_last_region,
                    ),
                    row.e_bar,
                    row.g_bar,
                    row.g_count,
                    row.spacing_mm,
                ),
            )
            for option_idx, result in enumerate(rows_sorted, start=1):
                transverse_arrangement = ""
                if result.e_bar and result.spacing_mm > 0:
                    if result.g_bar and result.g_count > 0:
                        transverse_arrangement = f"1E {result.e_bar} + {result.g_count}G {result.g_bar} @ {result.spacing_mm} mm"
                    else:
                        transverse_arrangement = f"1E {result.e_bar} @ {result.spacing_mm} mm"

                cantidad_estribos_region = 0
                if result.spacing_mm > 0 and region_length_mm > 0.0:
                    cantidad_estribos_region = _stirrup_count_for_reporting(
                        region_length_mm,
                        result.spacing_mm,
                        is_first_region=is_first_region,
                        is_last_region=is_last_region,
                    )

                peso_unitario_estribo_kg = float(result.stirrup_unit_weight_kg or 0.0)
                peso_transversal_region_kg = peso_unitario_estribo_kg * cantidad_estribos_region
                transverse_sheet.append(
                    [
                        beam_id,
                        span_id,
                        region_id,
                        option_idx,
                        "cumple",
                        transverse_arrangement,
                        result.controlling_limit,
                        result.spacing_mm,
                        cantidad_estribos_region,
                        peso_unitario_estribo_kg,
                        peso_transversal_region_kg,
                    ]
                )


    workbook.save(path)


def write_run_log(path: Path, lines: Iterable[str]) -> None:
    content = "\n".join(lines) + "\n"
    path.write_text(content, encoding="utf-8")


def span_summary_from_results(beam_id: str, span_id: str, results: list[RegionDesignResult], message: str = "") -> SpanSummary:
    ok_regions = sum(1 for result in results if result.status == "ok")
    fail_regions = len(results) - ok_regions
    return SpanSummary(
        beam_id=beam_id,
        span_id=span_id,
        total_regions=len(results),
        ok_regions=ok_regions,
        fail_regions=fail_regions,
        status="ok" if fail_regions == 0 else "fail",
        message=message,
    )


def beam_summary_from_spans(beam_id: str, span_summaries: list[SpanSummary]) -> BeamSummary:
    ok_spans = sum(1 for summary in span_summaries if summary.status == "ok")
    fail_spans = len(span_summaries) - ok_spans
    return BeamSummary(
        beam_id=beam_id,
        total_spans=len(span_summaries),
        ok_spans=ok_spans,
        fail_spans=fail_spans,
        status="ok" if fail_spans == 0 else "fail",
    )





