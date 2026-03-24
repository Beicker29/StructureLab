from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.core.errors import InvalidUploadError
from rc_shear_torsion.io import REQUIRED_COLUMNS, as_text, find_header


def _natural_name_key(value: str) -> tuple[int, Any]:
    text = value.strip()
    if text.isdigit():
        return (0, int(text))
    return (1, text)


def extract_unique_names_from_excel(excel_bytes: bytes, sheet_name: str, field_name: str) -> set[str]:
    try:
        workbook = load_workbook(filename=BytesIO(excel_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidUploadError(f"No fue posible leer '{field_name}' como archivo Excel valido") from exc

    if sheet_name not in workbook.sheetnames:
        raise InvalidUploadError(f"La hoja '{sheet_name}' no existe en '{field_name}'")

    worksheet = workbook[sheet_name]
    try:
        header_row_idx, header_map = find_header(worksheet.iter_rows(min_row=1, max_row=30, values_only=True))
    except ValueError as exc:
        raise InvalidUploadError(
            f"'{field_name}' no tiene encabezado ETABS valido en las primeras 30 filas"
        ) from exc
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing_columns:
        raise InvalidUploadError(
            f"'{field_name}' no contiene columnas requeridas: {', '.join(missing_columns)}"
        )

    names: set[str] = set()
    for row in worksheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None:
            continue
        unique_name = as_text(row[header_map["UniqueName"]])
        station_value = row[header_map["Station"]]
        if not unique_name or station_value in (None, ""):
            continue
        names.add(unique_name)

    if not names:
        raise InvalidUploadError(
            f"'{field_name}' no contiene filas validas con columnas UniqueName y Station"
        )
    return names


def parse_csv_names(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [item.strip() for item in raw.split(",") if item.strip()]


def parse_pairs_json(raw: str | None) -> list[dict[str, str]]:
    if not raw:
        return []
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise InvalidUploadError("frame_pairs_json debe ser un JSON valido") from exc

    if not isinstance(payload, list):
        raise InvalidUploadError("frame_pairs_json debe ser una lista de objetos")

    parsed_pairs: list[dict[str, str]] = []
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise InvalidUploadError(f"frame_pairs_json[{index}] debe ser un objeto")
        seismic = str(item.get("seismic", "")).strip()
        gravity = str(item.get("gravity", "")).strip()
        span_id = str(item.get("id", "")).strip() or f"S{index}"
        if not seismic or not gravity:
            raise InvalidUploadError(
                f"frame_pairs_json[{index}] requiere 'seismic' y 'gravity'"
            )
        parsed_pairs.append({"id": span_id, "seismic": seismic, "gravity": gravity})
    return parsed_pairs


def _parse_optional_float(value: Any, field_name: str) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise InvalidUploadError(f"{field_name} debe ser numerico") from exc


def _parse_optional_int(value: Any, field_name: str) -> int | None:
    if value in (None, ""):
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise InvalidUploadError(f"{field_name} debe ser entero") from exc
    return parsed


def _parse_confined_flag(value: Any) -> bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in {"1", "true", "t", "yes", "y", "si", "s", "confinado", "c"}:
        return True
    if text in {"0", "false", "f", "no", "n", "nc", "no_confinado", "noconfinado"}:
        return False
    return None


def _validate_regions_cover_span(regions: list[dict[str, Any]], span_id: str) -> None:
    if not regions:
        raise InvalidUploadError(f"span_layout_json[{span_id}] requiere al menos una region")
    ordered = sorted(regions, key=lambda item: float(item["from"]))
    tol = 1.0e-9
    if abs(float(ordered[0]["from"]) - 0.0) > tol:
        raise InvalidUploadError(
            f"span_layout_json[{span_id}] debe iniciar en from=0.0"
        )
    if abs(float(ordered[-1]["to"]) - 1.0) > tol:
        raise InvalidUploadError(
            f"span_layout_json[{span_id}] debe terminar en to=1.0"
        )
    current = 0.0
    for region in ordered:
        start = float(region["from"])
        end = float(region["to"])
        if abs(start - current) > tol:
            raise InvalidUploadError(
                f"span_layout_json[{span_id}] tiene gap/overlap antes de region '{region['id']}'"
            )
        if end <= start:
            raise InvalidUploadError(
                f"span_layout_json[{span_id}] region '{region['id']}' requiere from < to"
            )
        current = end


def parse_span_layout_json(raw: str | None) -> list[dict[str, Any]]:
    if not raw:
        return []
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise InvalidUploadError("span_layout_json debe ser un JSON valido") from exc
    if not isinstance(payload, list):
        raise InvalidUploadError("span_layout_json debe ser una lista de vanos")

    parsed_spans: list[dict[str, Any]] = []
    used_ids: set[str] = set()
    for span_index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise InvalidUploadError(f"span_layout_json[{span_index}] debe ser un objeto")

        span_id = str(item.get("id") or item.get("span_id") or f"S{span_index}").strip()
        if not span_id:
            raise InvalidUploadError(f"span_layout_json[{span_index}] requiere id")
        if span_id in used_ids:
            raise InvalidUploadError(f"span_layout_json tiene id repetido: '{span_id}'")
        used_ids.add(span_id)

        seismic = str(item.get("seismic", "")).strip()
        gravity = str(item.get("gravity", "")).strip()
        if not seismic or not gravity:
            raise InvalidUploadError(
                f"span_layout_json[{span_index}] requiere 'seismic' y 'gravity'"
            )

        ratio_raw = (
            item.get("c_ratio_extremos")
            if item.get("c_ratio_extremos") is not None
            else item.get("fraccion_c_extremos")
        )
        if ratio_raw is None:
            ratio_raw = item.get("region_c_ratio")
        c_ratio_extremos = _parse_optional_float(
            ratio_raw, f"span_layout_json[{span_index}].c_ratio_extremos"
        )
        if c_ratio_extremos is not None and not (0.0 < c_ratio_extremos < 0.5):
            raise InvalidUploadError(
                f"span_layout_json[{span_index}].c_ratio_extremos debe estar entre 0 y 0.5"
            )

        d_ratio_raw = (
            item.get("d_ratio")
            if item.get("d_ratio") is not None
            else item.get("fraccion_d")
        )
        d_ratio = _parse_optional_float(
            d_ratio_raw,
            f"span_layout_json[{span_index}].d_ratio",
        )
        if d_ratio is not None and not (0.0 < d_ratio <= 1.0):
            raise InvalidUploadError(
                f"span_layout_json[{span_index}].d_ratio debe estar en (0, 1]"
            )

        support_left_mm = _parse_optional_float(
            item.get("support_left_mm")
            if item.get("support_left_mm") is not None
            else item.get("apoyo_izq_mm"),
            f"span_layout_json[{span_index}].support_left_mm",
        )
        support_right_mm = _parse_optional_float(
            item.get("support_right_mm")
            if item.get("support_right_mm") is not None
            else item.get("apoyo_der_mm"),
            f"span_layout_json[{span_index}].support_right_mm",
        )
        if support_left_mm is None:
            raise InvalidUploadError(
                f"span_layout_json[{span_index}].support_left_mm es obligatorio"
            )
        if support_right_mm is None:
            raise InvalidUploadError(
                f"span_layout_json[{span_index}].support_right_mm es obligatorio"
            )
        if support_left_mm is not None and support_left_mm < 0.0:
            raise InvalidUploadError(
                f"span_layout_json[{span_index}].support_left_mm debe ser >= 0"
            )
        if support_right_mm is not None and support_right_mm < 0.0:
            raise InvalidUploadError(
                f"span_layout_json[{span_index}].support_right_mm debe ser >= 0"
            )

        raw_regions = item.get("regions")
        if raw_regions is None:
            raw_regions = item.get("regiones")
        regions: list[dict[str, Any]] = []
        if raw_regions is not None:
            if not isinstance(raw_regions, list):
                raise InvalidUploadError(
                    f"span_layout_json[{span_index}].regions debe ser una lista"
                )
            if not raw_regions:
                raise InvalidUploadError(
                    f"span_layout_json[{span_index}].regions no puede estar vacia"
                )
            for region_index, region_item in enumerate(raw_regions, start=1):
                if not isinstance(region_item, dict):
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}] debe ser un objeto"
                    )
                region_id = str(
                    region_item.get("id")
                    or region_item.get("region_id")
                    or f"R{region_index}"
                ).strip()
                if not region_id:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}] requiere id"
                    )

                start_raw = (
                    region_item.get("from")
                    if region_item.get("from") is not None
                    else region_item.get("desde")
                )
                end_raw = (
                    region_item.get("to")
                    if region_item.get("to") is not None
                    else region_item.get("hasta")
                )
                if start_raw in (None, "") or end_raw in (None, ""):
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}] requiere from y to"
                    )
                start_ratio = _parse_optional_float(
                    start_raw,
                    f"span_layout_json[{span_index}].regions[{region_index}].from",
                )
                end_ratio = _parse_optional_float(
                    end_raw,
                    f"span_layout_json[{span_index}].regions[{region_index}].to",
                )
                assert start_ratio is not None
                assert end_ratio is not None
                if not (0.0 <= start_ratio <= 1.0 and 0.0 <= end_ratio <= 1.0):
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}] debe estar dentro de [0,1]"
                    )
                if start_ratio >= end_ratio:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}] requiere from < to"
                    )

                region_type = str(region_item.get("type") or "").strip().upper()
                if region_type not in {"C", "NC"}:
                    confined = _parse_confined_flag(
                        region_item.get("is_confined")
                        if region_item.get("is_confined") is not None
                        else region_item.get("confinado")
                    )
                    if confined is None:
                        raise InvalidUploadError(
                            f"span_layout_json[{span_index}].regions[{region_index}] requiere type=C/NC o confinado=true/false"
                        )
                    region_type = "C" if confined else "NC"

                min_branches = _parse_optional_int(
                    region_item.get("min_branches")
                    if region_item.get("min_branches") is not None
                    else region_item.get("ramas_minimas"),
                    f"span_layout_json[{span_index}].regions[{region_index}].min_branches",
                )
                if min_branches is not None and min_branches < 1:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}].min_branches debe ser >= 1"
                    )

                spacing_mm = _parse_optional_int(
                    region_item.get("spacing_mm"),
                    f"span_layout_json[{span_index}].regions[{region_index}].spacing_mm",
                )
                if spacing_mm is not None and spacing_mm <= 0:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}].spacing_mm debe ser > 0"
                    )

                d_mm = _parse_optional_float(
                    region_item.get("d_mm"),
                    f"span_layout_json[{span_index}].regions[{region_index}].d_mm",
                )
                width_mm = _parse_optional_float(
                    region_item.get("width_mm"),
                    f"span_layout_json[{span_index}].regions[{region_index}].width_mm",
                )
                height_mm = _parse_optional_float(
                    region_item.get("height_mm"),
                    f"span_layout_json[{span_index}].regions[{region_index}].height_mm",
                )
                if d_mm is not None and d_mm <= 0:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}].d_mm debe ser > 0"
                    )
                if width_mm is not None and width_mm <= 0:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}].width_mm debe ser > 0"
                    )
                if height_mm is not None and height_mm <= 0:
                    raise InvalidUploadError(
                        f"span_layout_json[{span_index}].regions[{region_index}].height_mm debe ser > 0"
                    )

                db_bar = region_item.get("db_bar")
                db_bar_text = str(db_bar).strip() if db_bar not in (None, "") else None

                regions.append(
                    {
                        "id": region_id,
                        "from": float(start_ratio),
                        "to": float(end_ratio),
                        "type": region_type,
                        "min_branches": min_branches,
                        "spacing_mm": spacing_mm,
                        "d_mm": d_mm,
                        "db_bar": db_bar_text,
                        "width_mm": width_mm,
                        "height_mm": height_mm,
                    }
                )

            _validate_regions_cover_span(regions, span_id)

        parsed_spans.append(
            {
                "id": span_id,
                "seismic": seismic,
                "gravity": gravity,
                "c_ratio_extremos": c_ratio_extremos,
                "d_ratio": d_ratio,
                "support_left_mm": support_left_mm,
                "support_right_mm": support_right_mm,
                "regions": regions,
            }
        )

    return parsed_spans


def parse_optimization_overrides(raw: str | None) -> dict[str, Any] | None:
    if not raw:
        return None
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise InvalidUploadError("optimization_overrides_json debe ser JSON valido") from exc
    if not isinstance(payload, dict):
        raise InvalidUploadError("optimization_overrides_json debe ser un objeto JSON")
    return payload


def resolve_span_pairs(
    *,
    seismic_names: set[str],
    gravity_names: set[str],
    frame_names_csv: str | None,
    frame_pairs_json: str | None,
) -> list[dict[str, str]]:
    span_pairs = parse_pairs_json(frame_pairs_json)
    if span_pairs:
        for pair in span_pairs:
            if pair["seismic"] not in seismic_names:
                raise InvalidUploadError(f"El vano seismic '{pair['seismic']}' no existe en seismic_excel")
            if pair["gravity"] not in gravity_names:
                raise InvalidUploadError(f"El vano gravity '{pair['gravity']}' no existe en gravity_excel")
        return span_pairs

    csv_names = parse_csv_names(frame_names_csv)
    if csv_names:
        for frame_name in csv_names:
            if frame_name not in seismic_names:
                raise InvalidUploadError(f"El vano '{frame_name}' no existe en seismic_excel")
            if frame_name not in gravity_names:
                raise InvalidUploadError(f"El vano '{frame_name}' no existe en gravity_excel")
        return [
            {"id": f"S{index}", "seismic": frame_name, "gravity": frame_name}
            for index, frame_name in enumerate(csv_names, start=1)
        ]

    shared = sorted(seismic_names & gravity_names, key=_natural_name_key)
    if not shared:
        raise InvalidUploadError(
            "No se encontraron UniqueName comunes entre seismic_excel y gravity_excel. "
            "Usa frame_pairs_json para mapear vanos distintos."
        )
    return [
        {"id": f"S{index}", "seismic": frame_name, "gravity": frame_name}
        for index, frame_name in enumerate(shared, start=1)
    ]



def extract_design_sections_by_unique_name(
    excel_bytes: bytes,
    sheet_name: str,
    field_name: str,
) -> dict[str, str]:
    try:
        workbook = load_workbook(filename=BytesIO(excel_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidUploadError(f"No fue posible leer '{field_name}' como archivo Excel valido") from exc

    if sheet_name not in workbook.sheetnames:
        raise InvalidUploadError(f"La hoja '{sheet_name}' no existe en '{field_name}'")

    worksheet = workbook[sheet_name]
    try:
        header_row_idx, header_map = find_header(worksheet.iter_rows(min_row=1, max_row=30, values_only=True))
    except ValueError as exc:
        raise InvalidUploadError(
            f"'{field_name}' no tiene encabezado ETABS valido en las primeras 30 filas"
        ) from exc
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing_columns:
        raise InvalidUploadError(
            f"'{field_name}' no contiene columnas requeridas: {', '.join(missing_columns)}"
        )

    by_name: dict[str, dict[str, int]] = {}
    for row in worksheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None:
            continue
        unique_name = as_text(row[header_map["UniqueName"]])
        station_value = row[header_map["Station"]]
        if not unique_name or station_value in (None, ""):
            continue
        design_section = as_text(row[header_map["DesignSect"]])
        if not design_section:
            continue
        counts = by_name.setdefault(unique_name, {})
        counts[design_section] = counts.get(design_section, 0) + 1

    output: dict[str, str] = {}
    for unique_name, counts in by_name.items():
        if not counts:
            continue
        output[unique_name] = sorted(counts.items(), key=lambda item: (-item[1], item[0]))[0][0]
    return output


def _parse_positive_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if parsed <= 0.0:
        return None
    return parsed


def extract_geometry_sections(excel_bytes: bytes, field_name: str) -> dict[str, dict[str, float]]:
    try:
        workbook = load_workbook(filename=BytesIO(excel_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidUploadError(f"No fue posible leer '{field_name}' como archivo Excel valido") from exc

    geometry_sheet = None
    geometry_header_row_idx = 0
    geometry_header_map: dict[str, int] = {}

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=80, values_only=True),
            start=1,
        ):
            header_map = {as_text(value): i for i, value in enumerate(row) if as_text(value)}
            if {"Name", "Depth", "Width"}.issubset(header_map):
                geometry_sheet = worksheet
                geometry_header_row_idx = row_index
                geometry_header_map = header_map
                break
        if geometry_sheet is not None:
            break

    if geometry_sheet is None:
        raise InvalidUploadError(
            f"'{field_name}' no contiene una tabla con columnas Name, Depth y Width"
        )

    output: dict[str, dict[str, float]] = {}
    for row in geometry_sheet.iter_rows(min_row=geometry_header_row_idx + 1, values_only=True):
        if row is None:
            continue
        section_name = as_text(row[geometry_header_map["Name"]])
        if not section_name:
            continue
        depth = _parse_positive_float(row[geometry_header_map["Depth"]])
        width = _parse_positive_float(row[geometry_header_map["Width"]])
        if depth is None or width is None:
            continue

        existing = output.get(section_name)
        if existing is None:
            output[section_name] = {"height_mm": depth, "width_mm": width}
            continue

        # En secciones repetidas conservamos la geometria mas conservadora (maxima).
        output[section_name] = {
            "height_mm": max(existing["height_mm"], depth),
            "width_mm": max(existing["width_mm"], width),
        }

    if not output:
        raise InvalidUploadError(
            f"'{field_name}' no contiene filas validas para Name/Depth/Width"
        )
    return output
