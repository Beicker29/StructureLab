from __future__ import annotations

from collections import defaultdict
import re
from pathlib import Path
from threading import RLock
from typing import Any

from openpyxl import load_workbook

from app.services.job_service import get_job, get_job_case_payload
from rc_shear_torsion.design import bar_mass_kg_per_m

_PREVIEW_CACHE_LOCK = RLock()
_PREVIEW_CACHE: dict[str, tuple[tuple[Any, ...], dict[str, Any]]] = {}



def _artifact_mtime(path: Path | None) -> float | None:
    if path is None or not path.exists():
        return None
    try:
        return float(path.stat().st_mtime)
    except OSError:
        return None

def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    return {_normalize_header(value): index for index, value in enumerate(row)}


def _find_col(columns: dict[str, int], *names: str) -> int | None:
    for name in names:
        key = _normalize_header(name)
        if key in columns:
            return columns[key]
    return None


def _as_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _as_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if parsed <= 0:
        return None
    return parsed


def _as_non_negative_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if parsed < 0:
        return None
    return parsed


def _as_int(value: Any) -> int | None:
    parsed = _as_float(value)
    if parsed is None:
        return None
    return int(round(parsed))


def _as_non_negative_int(value: Any) -> int | None:
    parsed = _as_non_negative_float(value)
    if parsed is None:
        return None
    return int(round(parsed))


def _longitudinal_weight_from_arrangement(bar: str, count: int | None, length_mm: float | None) -> float:
    if not bar or count is None or count <= 0:
        return 0.0
    if length_mm is None or length_mm <= 0.0:
        return 0.0
    return bar_mass_kg_per_m(bar, int(count)) * (float(length_mm) / 1000.0)

def _norm_match(value: str) -> str:
    return " ".join(value.lower().split())

def _long_arrangement_label(bar: str, count: int | None, *, empty_label: str = "") -> str:
    if not bar or count is None or count <= 0:
        return empty_label
    return f"{count} x {bar}"


def _parse_long_arrangement_label(label: str) -> tuple[str, int | None]:
    text = _as_text(label)
    if not text:
        return "", None
    match = re.match(r"^(\d+)\s*x\s*(#\d+)$", text, re.IGNORECASE)
    if not match:
        return "", None
    return match.group(2).upper(), int(match.group(1))


def _resolve_region_longitudinal_weight_kg(option_row: dict[str, Any], region_length_mm: float | None) -> float:
    explicit = _as_non_negative_float(option_row.get("weight_longitudinal_kg"))
    if explicit is not None:
        return float(explicit)

    base_weight = _as_non_negative_float(option_row.get("weight_base_kg"))
    if base_weight is None:
        base_weight = _longitudinal_weight_from_arrangement(
            _as_text(option_row.get("base_long_bar")),
            _as_int(option_row.get("base_long_count")),
            region_length_mm,
        )

    additional_weight = _as_non_negative_float(option_row.get("weight_additional_kg"))
    if additional_weight is None:
        additional_weight = _longitudinal_weight_from_arrangement(
            _as_text(option_row.get("extra_long_bar")),
            _as_int(option_row.get("extra_long_count")),
            region_length_mm,
        )

    return float(base_weight or 0.0) + float(additional_weight or 0.0)


def _resolve_transverse_weight_kg(
    *,
    stirrup_count: int | None,
    stirrup_unit_weight_kg: float | None,
    explicit_weight_kg: float | None,
) -> float | None:
    if stirrup_count is not None and stirrup_unit_weight_kg is not None:
        return float(stirrup_count) * float(stirrup_unit_weight_kg)
    if explicit_weight_kg is not None:
        return float(explicit_weight_kg)
    return None


def _read_optimized_regions(path: Path | None) -> dict[tuple[str, str], dict[str, Any]]:
    if path is None or not path.exists():
        return {}

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["optimized_regions"] if "optimized_regions" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    if header_row is None:
        return {}

    columns = _header_map(header_row)
    span_col = _find_col(columns, "span_id", "vano_id")
    region_col = _find_col(columns, "region_id")
    e_bar_col = _find_col(columns, "e_bar")
    g_bar_col = _find_col(columns, "g_bar")
    g_count_col = _find_col(columns, "g_count")
    spacing_col = _find_col(columns, "spacing_mm")
    long_bar_col = _find_col(columns, "long_bar")
    long_count_col = _find_col(columns, "long_count")
    longitudinal_arrangement_col = _find_col(columns, "longitudinal_arrangement")
    base_long_bar_col = _find_col(columns, "base_long_bar")
    base_long_count_col = _find_col(columns, "base_long_count")
    extra_long_bar_col = _find_col(columns, "extra_long_bar")
    extra_long_count_col = _find_col(columns, "extra_long_count")
    longitudinal_mode_col = _find_col(columns, "longitudinal_mode")
    is_deep_beam_col = _find_col(columns, "is_deep_beam")
    status_col = _find_col(columns, "status")
    if span_col is None or region_col is None:
        return {}

    output: dict[tuple[str, str], dict[str, Any]] = {}
    for row in iterator:
        span_id = _as_text(row[span_col])
        region_id = _as_text(row[region_col])
        if not span_id or not region_id:
            continue
        status = _as_text(row[status_col]) if status_col is not None else "ok"
        if status and status.lower() not in {"ok", "cumple"}:
            continue

        e_bar = _as_text(row[e_bar_col]) if e_bar_col is not None else ""
        g_bar = _as_text(row[g_bar_col]) if g_bar_col is not None else ""
        g_count = _as_int(row[g_count_col]) if g_count_col is not None else None
        spacing = _as_int(row[spacing_col]) if spacing_col is not None else None
        long_bar = _as_text(row[long_bar_col]) if long_bar_col is not None else ""
        long_count = _as_int(row[long_count_col]) if long_count_col is not None else None
        longitudinal_arrangement = _as_text(row[longitudinal_arrangement_col]) if longitudinal_arrangement_col is not None else ""
        base_long_bar = _as_text(row[base_long_bar_col]) if base_long_bar_col is not None else ""
        base_long_count = _as_int(row[base_long_count_col]) if base_long_count_col is not None else None
        extra_long_bar = _as_text(row[extra_long_bar_col]) if extra_long_bar_col is not None else ""
        extra_long_count = _as_int(row[extra_long_count_col]) if extra_long_count_col is not None else None
        longitudinal_mode = _as_text(row[longitudinal_mode_col]) if longitudinal_mode_col is not None else ""
        is_deep_beam = _as_text(row[is_deep_beam_col]) if is_deep_beam_col is not None else ""

        if e_bar and g_bar and g_count is not None and spacing is not None:
            transverse = f"1E {e_bar} + {g_count}G {g_bar} @ {spacing} mm"
        elif e_bar and spacing is not None:
            transverse = f"1E {e_bar} @ {spacing} mm"
        else:
            transverse = ""

        if longitudinal_arrangement:
            longitudinal_schedule = longitudinal_arrangement
            longitudinal = longitudinal_arrangement.replace(" x ", "")
        elif long_count is not None and long_bar:
            longitudinal = f"{long_count}{long_bar}"
            longitudinal_schedule = f"{long_count} x {long_bar}"
        else:
            longitudinal = "no se requiere"
            longitudinal_schedule = "no se requiere"

        output[(span_id, region_id)] = {
            "spacing_mm": spacing,
            "transverse_label": transverse,
            "longitudinal_label": longitudinal,
            "longitudinal_schedule_label": longitudinal_schedule,
            "base_long_bar": base_long_bar,
            "base_long_count": base_long_count,
            "extra_long_bar": extra_long_bar,
            "extra_long_count": extra_long_count,
            "longitudinal_mode": longitudinal_mode,
            "is_deep_beam": is_deep_beam,
        }
    return output


def _read_schedule_rows(path: Path | None) -> dict[tuple[str, str], list[dict[str, Any]]]:
    if path is None or not path.exists():
        return {}

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["por_region"] if "por_region" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    if header_row is None:
        return {}

    columns = _header_map(header_row)
    span_col = _find_col(columns, "span_id", "vano_id")
    region_col = _find_col(columns, "region_id")
    option_col = _find_col(columns, "option", "opcion")
    length_col = _find_col(columns, "longitud_region_mm", "region_length_mm", "length_mm")
    trans_col = _find_col(columns, "arreglo_transversal", "transverse_arrangement")
    long_col = _find_col(columns, "arreglo_longitudinal", "longitudinal_arrangement")
    base_arr_col = _find_col(columns, "arreglo_longitudinal_base", "base_longitudinal_arrangement")
    add_arr_col = _find_col(columns, "arreglo_longitudinal_adicional", "additional_longitudinal_arrangement")
    base_long_bar_col = _find_col(columns, "base_long_bar")
    base_long_count_col = _find_col(columns, "base_long_count")
    extra_long_bar_col = _find_col(columns, "extra_long_bar")
    extra_long_count_col = _find_col(columns, "extra_long_count")
    stirrup_count_col = _find_col(columns, "cantidad_estribos_region", "stirrup_count_region", "stirrup_count")
    stirrup_unit_weight_col = _find_col(columns, "peso_unitario_estribo_kg", "stirrup_unit_weight_kg")
    weight_trans_col = _find_col(columns, "peso_transversal_region_kg", "transverse_weight_region_kg")
    weight_long_col = _find_col(columns, "peso_longitudinal_region_kg", "longitudinal_weight_region_kg")
    weight_total_col = _find_col(columns, "peso_total_region_kg", "total_weight_region_kg")
    status_col = _find_col(columns, "estado", "status")
    if span_col is None or region_col is None or length_col is None:
        return {}

    output: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in iterator:
        span_id = _as_text(row[span_col])
        region_id = _as_text(row[region_col])
        if not span_id or not region_id:
            continue
        length_mm = _as_float(row[length_col])
        if length_mm is None:
            continue
        status = _as_text(row[status_col]).lower() if status_col is not None else ""
        if status and status not in {"cumple", "ok"}:
            continue
        option = _as_int(row[option_col]) if option_col is not None else None
        transverse = _as_text(row[trans_col]) if trans_col is not None else ""
        longitudinal = _as_text(row[long_col]) if long_col is not None else ""
        base_long_bar = _as_text(row[base_long_bar_col]) if base_long_bar_col is not None else ""
        base_long_count = _as_int(row[base_long_count_col]) if base_long_count_col is not None else None
        extra_long_bar = _as_text(row[extra_long_bar_col]) if extra_long_bar_col is not None else ""
        extra_long_count = _as_int(row[extra_long_count_col]) if extra_long_count_col is not None else None
        base_longitudinal_label = (
            _as_text(row[base_arr_col])
            if base_arr_col is not None
            else _long_arrangement_label(base_long_bar, base_long_count, empty_label="no se requiere")
        )
        additional_longitudinal_label = _as_text(row[add_arr_col]) if add_arr_col is not None else ""
        if not additional_longitudinal_label:
            additional_longitudinal_label = _long_arrangement_label(extra_long_bar, extra_long_count, empty_label="no se requiere")

        if (not base_long_bar or base_long_count is None) and base_longitudinal_label:
            parsed_bar, parsed_count = _parse_long_arrangement_label(base_longitudinal_label)
            if parsed_bar and parsed_count is not None:
                base_long_bar = parsed_bar
                base_long_count = parsed_count

        if (not extra_long_bar or extra_long_count is None) and additional_longitudinal_label:
            parsed_bar, parsed_count = _parse_long_arrangement_label(additional_longitudinal_label)
            if parsed_bar and parsed_count is not None:
                extra_long_bar = parsed_bar
                extra_long_count = parsed_count

        stirrup_count = _as_non_negative_int(row[stirrup_count_col]) if stirrup_count_col is not None else None
        stirrup_unit_weight_kg = (
            _as_non_negative_float(row[stirrup_unit_weight_col]) if stirrup_unit_weight_col is not None else None
        )
        explicit_transverse_weight_kg = _as_non_negative_float(row[weight_trans_col]) if weight_trans_col is not None else None
        weight_trans_kg = _resolve_transverse_weight_kg(
            stirrup_count=stirrup_count,
            stirrup_unit_weight_kg=stirrup_unit_weight_kg,
            explicit_weight_kg=explicit_transverse_weight_kg,
        )
        weight_long_kg = _as_non_negative_float(row[weight_long_col]) if weight_long_col is not None else None
        weight_total_kg = _as_non_negative_float(row[weight_total_col]) if weight_total_col is not None else None
        weight_base_kg = _longitudinal_weight_from_arrangement(base_long_bar, base_long_count, length_mm)
        weight_additional_kg = _longitudinal_weight_from_arrangement(extra_long_bar, extra_long_count, length_mm)
        output[(span_id, region_id)].append(
            {
                "length_mm": length_mm,
                "option": option,
                "transverse_label": transverse,
                "longitudinal_label": longitudinal,
                "base_longitudinal_label": base_longitudinal_label,
                "additional_longitudinal_label": additional_longitudinal_label,
                "base_long_bar": base_long_bar,
                "base_long_count": base_long_count,
                "extra_long_bar": extra_long_bar,
                "extra_long_count": extra_long_count,
                "stirrup_count": stirrup_count,
                "stirrup_unit_weight_kg": stirrup_unit_weight_kg,
                "weight_transverse_kg": weight_trans_kg,
                "weight_longitudinal_kg": weight_long_kg,
                "weight_base_kg": weight_base_kg,
                "weight_additional_kg": weight_additional_kg,
                "weight_total_kg": weight_total_kg,
            }
        )
    return output


def _read_transverse_options(path: Path | None) -> dict[tuple[str, str], list[dict[str, Any]]]:
    if path is None or not path.exists():
        return {}

    workbook = load_workbook(path, data_only=True, read_only=True)
    if "transversales" not in workbook.sheetnames:
        return {}

    sheet = workbook["transversales"]
    iterator = sheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    if header_row is None:
        return {}

    columns = _header_map(header_row)
    span_col = _find_col(columns, "vano_id", "span_id")
    region_col = _find_col(columns, "region_id")
    status_col = _find_col(columns, "estado", "status")
    label_col = _find_col(columns, "arreglo_transversal", "transverse_arrangement")
    stirrup_count_col = _find_col(columns, "cantidad_estribos_region", "stirrup_count_region", "stirrup_count")
    stirrup_unit_weight_col = _find_col(columns, "peso_unitario_estribo_kg", "stirrup_unit_weight_kg")
    weight_col = _find_col(columns, "peso_transversal_region_kg", "transverse_weight_region_kg")

    if span_col is None or region_col is None or label_col is None:
        return {}

    output: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in iterator:
        span_id = _as_text(row[span_col])
        region_id = _as_text(row[region_col])
        if not span_id or not region_id:
            continue

        status = _as_text(row[status_col]).lower() if status_col is not None else ""
        if status and status not in {"cumple", "ok"}:
            continue

        label = _as_text(row[label_col])
        if not label:
            continue

        stirrup_count = _as_non_negative_int(row[stirrup_count_col]) if stirrup_count_col is not None else None
        stirrup_unit_weight_kg = (
            _as_non_negative_float(row[stirrup_unit_weight_col]) if stirrup_unit_weight_col is not None else None
        )
        explicit_weight = _as_non_negative_float(row[weight_col]) if weight_col is not None else None
        weight = _resolve_transverse_weight_kg(
            stirrup_count=stirrup_count,
            stirrup_unit_weight_kg=stirrup_unit_weight_kg,
            explicit_weight_kg=explicit_weight,
        )

        output[(span_id, region_id)].append(
            {
                "label": label,
                "weight_kg": float(weight) if weight is not None else None,
                "stirrup_count": stirrup_count,
                "stirrup_unit_weight_kg": (
                    float(stirrup_unit_weight_kg) if stirrup_unit_weight_kg is not None else None
                ),
            }
        )

    for key, rows in output.items():
        best_by_label: dict[str, dict[str, Any]] = {}
        for item in rows:
            label = _as_text(item.get("label"))
            if not label:
                continue
            sort_weight = (
                float(item["weight_kg"])
                if item.get("weight_kg") is not None
                else float("inf")
            )
            current = best_by_label.get(label)
            if current is None or sort_weight < current["_sort_weight"]:
                best_by_label[label] = {
                    "label": label,
                    "weight_kg": item.get("weight_kg"),
                    "stirrup_count": item.get("stirrup_count"),
                    "stirrup_unit_weight_kg": item.get("stirrup_unit_weight_kg"),
                    "_sort_weight": sort_weight,
                }

        ordered = sorted(
            best_by_label.values(),
            key=lambda item: (item["_sort_weight"], item["label"]),
        )[:10]
        output[key] = [
            {
                "label": item["label"],
                "weight_kg": item.get("weight_kg"),
                "stirrup_count": item.get("stirrup_count"),
                "stirrup_unit_weight_kg": item.get("stirrup_unit_weight_kg"),
            }
            for item in ordered
        ]

    return output

def _pick_schedule_row(
    rows: list[dict[str, Any]],
    optimized: dict[str, Any] | None,
) -> dict[str, Any] | None:
    if not rows:
        return None
    if optimized:
        opt_trans = _norm_match(str(optimized.get("transverse_label") or ""))
        opt_long = _norm_match(str(optimized.get("longitudinal_schedule_label") or ""))
        if opt_trans:
            for row in rows:
                row_trans = _norm_match(str(row.get("transverse_label") or ""))
                row_long = _norm_match(str(row.get("longitudinal_label") or ""))
                if row_trans == opt_trans and (not opt_long or row_long == opt_long):
                    return row
    rows_sorted = sorted(
        rows,
        key=lambda item: (
            item.get("weight_total_kg") if item.get("weight_total_kg") is not None else float("inf"),
            item.get("option") if item.get("option") is not None else 999_999,
        ),
    )
    return rows_sorted[0] if rows_sorted else None


def _build_component_options(
    rows: list[dict[str, Any]],
    *,
    label_key: str,
    weight_key: str,
    max_items: int | None = 10,
) -> list[dict[str, Any]]:
    best_by_label: dict[str, dict[str, Any]] = {}
    for row in rows:
        label = _as_text(row.get(label_key))
        if not label:
            continue
        weight = _as_non_negative_float(row.get(weight_key))
        sort_weight = weight if weight is not None else float("inf")
        current = best_by_label.get(label)
        if current is None or sort_weight < current["_sort_weight"]:
            best_by_label[label] = {
                "label": label,
                "weight_kg": float(weight) if weight is not None else None,
                "_sort_weight": sort_weight,
            }
    ordered = sorted(best_by_label.values(), key=lambda item: (item["_sort_weight"], item["label"]))
    if max_items is not None:
        ordered = ordered[: max(1, int(max_items))]
    return [
        {
            "label": item["label"],
            "weight_kg": item["weight_kg"],
        }
        for item in ordered
    ]


def _build_transverse_component_options(
    rows: list[dict[str, Any]],
    *,
    max_items: int = 10,
) -> list[dict[str, Any]]:
    best_by_label: dict[str, dict[str, Any]] = {}
    for row in rows:
        label = _as_text(row.get("transverse_label"))
        if not label:
            continue

        stirrup_count = _as_non_negative_int(row.get("stirrup_count"))
        stirrup_unit_weight_kg = _as_non_negative_float(row.get("stirrup_unit_weight_kg"))
        explicit_weight = _as_non_negative_float(row.get("weight_transverse_kg"))
        weight = _resolve_transverse_weight_kg(
            stirrup_count=stirrup_count,
            stirrup_unit_weight_kg=stirrup_unit_weight_kg,
            explicit_weight_kg=explicit_weight,
        )
        sort_weight = weight if weight is not None else float("inf")
        current = best_by_label.get(label)
        if current is None or sort_weight < current["_sort_weight"]:
            best_by_label[label] = {
                "label": label,
                "weight_kg": float(weight) if weight is not None else None,
                "stirrup_count": stirrup_count,
                "stirrup_unit_weight_kg": (
                    float(stirrup_unit_weight_kg) if stirrup_unit_weight_kg is not None else None
                ),
                "_sort_weight": sort_weight,
            }

    ordered = sorted(best_by_label.values(), key=lambda item: (item["_sort_weight"], item["label"]))
    return [
        {
            "label": item["label"],
            "weight_kg": item["weight_kg"],
            "stirrup_count": item.get("stirrup_count"),
            "stirrup_unit_weight_kg": item.get("stirrup_unit_weight_kg"),
        }
        for item in ordered[: max(1, int(max_items))]
    ]

def _build_span_option_choices(region_rows: list[dict[str, Any]], max_items: int = 10) -> list[dict[str, Any]]:
    if not region_rows:
        return []

    coupled_rows = [
        row
        for row in region_rows
        if _as_text(row.get("longitudinal_mode")).lower() == "span_coupled"
        and isinstance(row.get("options"), list)
        and row.get("options")
    ]
    if len(coupled_rows) != len(region_rows):
        return []

    common_options: set[int] | None = None
    options_by_region: list[dict[int, dict[str, Any]]] = []
    for row in coupled_rows:
        by_option: dict[int, dict[str, Any]] = {}
        for option_row in row.get("options") or []:
            option_value = _as_int(option_row.get("option"))
            if option_value is None:
                continue
            current = by_option.get(option_value)
            current_weight = _as_non_negative_float((current or {}).get("weight_total_kg"))
            option_weight = _as_non_negative_float(option_row.get("weight_total_kg"))
            current_sort = current_weight if current_weight is not None else float("inf")
            option_sort = option_weight if option_weight is not None else float("inf")
            if current is None or option_sort < current_sort:
                by_option[option_value] = option_row
        if not by_option:
            return []
        options_by_region.append(by_option)
        region_options = set(by_option.keys())
        common_options = region_options if common_options is None else (common_options & region_options)

    if not common_options:
        return []

    output: list[dict[str, Any]] = []
    for option in sorted(common_options):
        total_weight = 0.0
        longitudinal_weight = 0.0
        for by_option in options_by_region:
            row = by_option[option]
            total_weight += float(row.get("weight_total_kg") or 0.0)
            longitudinal_weight += float(row.get("weight_longitudinal_kg") or 0.0)
        output.append(
            {
                "option": option,
                "total_weight_kg": total_weight,
                "long_weight_kg": longitudinal_weight,
                "regions": len(region_rows),
            }
        )

    output.sort(key=lambda item: (item["total_weight_kg"], item["option"]))
    return output[:max_items]


def _expand_span_coupled_options(
    *,
    longitudinal_rows: list[dict[str, Any]],
    transverse_options: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    if not longitudinal_rows:
        return []
    if not transverse_options:
        return longitudinal_rows

    long_sorted = sorted(
        longitudinal_rows,
        key=lambda item: (
            item.get("weight_total_kg") if item.get("weight_total_kg") is not None else float("inf"),
            item.get("option") if item.get("option") is not None else 999_999,
        ),
    )
    trans_sorted = sorted(
        transverse_options,
        key=lambda item: (
            _as_non_negative_float(item.get("weight_kg"))
            if _as_non_negative_float(item.get("weight_kg")) is not None
            else float("inf"),
            _as_text(item.get("label")),
        ),
    )

    expanded: list[dict[str, Any]] = []
    for long_row in long_sorted:
        long_weight = _resolve_region_longitudinal_weight_kg(
            long_row,
            _as_float(long_row.get("length_mm")),
        )
        long_option = _as_int(long_row.get("option"))
        for trans_rank, trans in enumerate(trans_sorted, start=1):
            label = _as_text(trans.get("label"))
            if not label:
                continue
            trans_weight = float(_as_non_negative_float(trans.get("weight_kg")) or 0.0)
            item = dict(long_row)
            item["option"] = long_option
            item["longitudinal_option"] = long_option
            item["transverse_option_rank"] = trans_rank
            item["transverse_label"] = label
            item["stirrup_count"] = _as_non_negative_int(trans.get("stirrup_count"))
            item["stirrup_unit_weight_kg"] = _as_non_negative_float(trans.get("stirrup_unit_weight_kg"))
            item["weight_transverse_kg"] = trans_weight
            item["weight_longitudinal_kg"] = long_weight
            item["weight_total_kg"] = trans_weight + long_weight
            expanded.append(item)

    if not expanded:
        return long_sorted
    expanded.sort(
        key=lambda item: (
            item.get("weight_total_kg") if item.get("weight_total_kg") is not None else float("inf"),
            item.get("option") if item.get("option") is not None else 999_999,
            item.get("transverse_option_rank") if item.get("transverse_option_rank") is not None else 999_999,
        )
    )
    return expanded


def _build_span_longitudinal_base_options(
    region_rows: list[dict[str, Any]],
    max_items: int = 10,
) -> list[dict[str, Any]]:
    if not region_rows:
        return []

    coupled_rows = [
        row
        for row in region_rows
        if _as_text(row.get("longitudinal_mode")).lower() == "span_coupled"
        and isinstance(row.get("options"), list)
        and row.get("options")
    ]
    if len(coupled_rows) != len(region_rows):
        return []

    common_base_labels: set[str] | None = None
    best_by_base_per_region: list[dict[str, dict[str, Any]]] = []
    for row in coupled_rows:
        region_length_mm = _as_float(row.get("length_mm"))
        best_by_base: dict[str, dict[str, Any]] = {}
        for option_row in row.get("options") or []:
            base_label = _as_text(option_row.get("base_longitudinal_label")) or "no se requiere"
            option_value = _as_int(option_row.get("option"))
            long_total_weight = _resolve_region_longitudinal_weight_kg(option_row, region_length_mm)

            base_weight = _as_non_negative_float(option_row.get("weight_base_kg"))
            if base_weight is None:
                base_weight = _longitudinal_weight_from_arrangement(
                    _as_text(option_row.get("base_long_bar")),
                    _as_int(option_row.get("base_long_count")),
                    region_length_mm,
                )

            item = {
                "base_label": base_label,
                "long_total_weight_kg": long_total_weight,
                "base_weight_region_kg": float(base_weight or 0.0),
                "option": option_value,
            }
            current = best_by_base.get(base_label)
            if current is None or (item["long_total_weight_kg"], item.get("option") or 999_999) < (
                current["long_total_weight_kg"],
                current.get("option") or 999_999,
            ):
                best_by_base[base_label] = item

        if not best_by_base:
            return []

        best_by_base_per_region.append(best_by_base)
        labels = set(best_by_base.keys())
        common_base_labels = labels if common_base_labels is None else (common_base_labels & labels)

    if not common_base_labels:
        return []

    options: list[dict[str, Any]] = []
    for base_label in sorted(common_base_labels):
        total_weight = 0.0
        long_weight = 0.0
        option_candidates: list[int] = []
        for by_base in best_by_base_per_region:
            item = by_base[base_label]
            total_weight += float(item["long_total_weight_kg"])
            long_weight += float(item["base_weight_region_kg"])
            option_value = item.get("option")
            if option_value is not None:
                option_candidates.append(int(option_value))

        option = min(option_candidates) if option_candidates else None
        slug = re.sub(r"[^a-z0-9]+", "_", _norm_match(base_label)).strip("_") or "base"
        options.append(
            {
                "value": f"base_{slug}",
                "option": option,
                "base_label": base_label,
                "long_weight_kg": long_weight,
                "total_weight_kg": total_weight,
            }
        )

    options.sort(
        key=lambda item: (
            item["total_weight_kg"],
            item["long_weight_kg"],
            item["base_label"],
        )
    )
    return options[:max_items]

def _build_span_longitudinal_option_sets(
    region_rows: list[dict[str, Any]],
    base_options: list[dict[str, Any]],
    max_items: int = 10,
) -> list[dict[str, Any]]:
    if not region_rows or not base_options:
        return []

    coupled_rows = [
        row
        for row in region_rows
        if _as_text(row.get("longitudinal_mode")).lower() == "span_coupled"
    ]
    if len(coupled_rows) != len(region_rows):
        return []

    sets: list[dict[str, Any]] = []
    for base in base_options:
        base_value = _as_text(base.get("value"))
        base_label = _as_text(base.get("base_label")) or "no se requiere"
        if not base_value:
            continue

        regions_payload: list[dict[str, Any]] = []
        total_weight = 0.0
        total_transverse = 0.0
        total_longitudinal = 0.0
        total_base = 0.0
        total_additional = 0.0
        feasible = True
        option_candidates: list[int] = []

        for row in coupled_rows:
            region_id = _as_text(row.get("region_id"))
            region_length_mm = _as_float(row.get("length_mm"))
            options = row.get("options") if isinstance(row.get("options"), list) else []
            base_rows = [
                option_row
                for option_row in options
                if (_as_text(option_row.get("base_longitudinal_label")) or "no se requiere") == base_label
            ]
            if not base_rows:
                feasible = False
                break

            best_row = sorted(
                base_rows,
                key=lambda item: (
                    _as_non_negative_float(item.get("weight_total_kg"))
                    if _as_non_negative_float(item.get("weight_total_kg")) is not None
                    else float("inf"),
                    _as_non_negative_float(item.get("weight_transverse_kg"))
                    if _as_non_negative_float(item.get("weight_transverse_kg")) is not None
                    else float("inf"),
                    _as_non_negative_float(item.get("weight_additional_kg"))
                    if _as_non_negative_float(item.get("weight_additional_kg")) is not None
                    else float("inf"),
                    _as_int(item.get("option")) if _as_int(item.get("option")) is not None else 999_999,
                ),
            )[0]

            trans_w = float(_as_non_negative_float(best_row.get("weight_transverse_kg")) or 0.0)
            long_w = float(_as_non_negative_float(best_row.get("weight_longitudinal_kg")) or 0.0)
            if long_w <= 0.0:
                long_w = _resolve_region_longitudinal_weight_kg(best_row, region_length_mm)
            total_w = _as_non_negative_float(best_row.get("weight_total_kg"))
            if total_w is None:
                total_w = trans_w + long_w
            else:
                total_w = float(total_w)

            base_w = _as_non_negative_float(best_row.get("weight_base_kg"))
            if base_w is None:
                base_w = _longitudinal_weight_from_arrangement(
                    _as_text(best_row.get("base_long_bar")),
                    _as_int(best_row.get("base_long_count")),
                    region_length_mm,
                )
            add_w = _as_non_negative_float(best_row.get("weight_additional_kg"))
            if add_w is None:
                add_w = max(float(long_w) - float(base_w or 0.0), 0.0)
            option_value = _as_int(best_row.get("option"))
            if option_value is not None:
                option_candidates.append(option_value)

            total_weight += float(total_w)
            total_transverse += float(trans_w)
            total_longitudinal += long_w
            total_base += float(base_w or 0.0)
            total_additional += float(add_w or 0.0)
            regions_payload.append(
                {
                    "region_id": region_id,
                    "transverse_label": _as_text(best_row.get("transverse_label")),
                    "additional_label": _as_text(best_row.get("additional_longitudinal_label")) or "no se requiere",
                    "longitudinal_label": _as_text(best_row.get("longitudinal_label")) or "no se requiere",
                    "weight_total_kg": float(total_w),
                    "weight_transverse_kg": float(trans_w),
                    "weight_longitudinal_kg": float(long_w),
                    "weight_base_region_kg": float(base_w or 0.0),
                    "weight_additional_kg": float(add_w or 0.0),
                    "option": option_value,
                }
            )

        if not feasible or not regions_payload:
            continue

        option_value = min(option_candidates) if option_candidates else _as_int(base.get("option"))
        sets.append(
            {
                "value": f"longset_{base_value}",
                "option": option_value,
                "base_label": base_label,
                "base_value": base_value,
                "total_weight_kg": float(total_weight),
                "total_transverse_weight_kg": float(total_transverse),
                "total_longitudinal_weight_kg": float(total_longitudinal),
                "total_base_weight_kg": float(total_base),
                "total_additional_weight_kg": float(total_additional),
                "regions": regions_payload,
            }
        )

    sets.sort(
        key=lambda item: (
            item["total_longitudinal_weight_kg"],
            item.get("total_weight_kg") if item.get("total_weight_kg") is not None else float("inf"),
            item.get("option") if item.get("option") is not None else 999_999,
            item.get("base_label") or "",
        )
    )
    for index, item in enumerate(sets, start=1):
        item["rank"] = index

    return sets[:max_items]


def _resolve_span_default_selection(span_row: dict[str, Any]) -> tuple[dict[str, Any], float]:
    span_id = _as_text(span_row.get("span_id")) or "S1"
    regions = span_row.get("regions") if isinstance(span_row.get("regions"), list) else []
    span_total_weight = 0.0
    span_default: dict[str, Any] = {
        "span_id": span_id,
        "longitudinal_mode": (
            _as_text((regions[0] if regions else {}).get("longitudinal_mode")) if regions else ""
        ),
        "span_option_set_value": None,
        "base_value": None,
        "base_label": None,
        "regions": [],
    }
    if not regions:
        return span_default, span_total_weight

    is_span_coupled = any(_as_text(region.get("longitudinal_mode")).lower() == "span_coupled" for region in regions)
    span_sets = (
        span_row.get("span_longitudinal_option_sets")
        if isinstance(span_row.get("span_longitudinal_option_sets"), list)
        else []
    )
    base_options = (
        span_row.get("longitudinal_base_options")
        if isinstance(span_row.get("longitudinal_base_options"), list)
        else []
    )
    best_set_by_total = (
        min(
            span_sets,
            key=lambda item: (
                _as_non_negative_float(item.get("total_weight_kg"))
                if _as_non_negative_float(item.get("total_weight_kg")) is not None
                else float("inf"),
                _as_non_negative_float(item.get("total_longitudinal_weight_kg"))
                if _as_non_negative_float(item.get("total_longitudinal_weight_kg")) is not None
                else float("inf"),
            ),
        )
        if is_span_coupled and span_sets
        else None
    )
    set_region_map: dict[str, dict[str, Any]] = {}
    if isinstance(best_set_by_total, dict):
        for item in best_set_by_total.get("regions") or []:
            region_id = _as_text((item or {}).get("region_id"))
            if not region_id:
                continue
            set_region_map[region_id] = item or {}
        span_default["span_option_set_value"] = _as_text(best_set_by_total.get("value")) or None
        span_default["base_label"] = _as_text(best_set_by_total.get("base_label")) or None
        span_default["base_value"] = _as_text(best_set_by_total.get("base_value")) or None
        if not span_default["base_value"] and span_default["base_label"]:
            matched_base = next(
                (
                    item
                    for item in base_options
                    if _as_text(item.get("base_label")) == _as_text(span_default.get("base_label"))
                ),
                None,
            )
            span_default["base_value"] = _as_text((matched_base or {}).get("value")) or None

    for region in regions:
        region_id = _as_text(region.get("region_id"))
        options = region.get("options") if isinstance(region.get("options"), list) else []
        if not region_id or not options:
            continue

        selected_row: dict[str, Any] | None = None
        if set_region_map:
            set_row = set_region_map.get(region_id) or {}
            selected_row = next(
                (
                    row
                    for row in options
                    if _norm_match(_as_text(row.get("transverse_label")))
                    == _norm_match(_as_text(set_row.get("transverse_label")))
                    and _norm_match(_as_text(row.get("longitudinal_label")))
                    == _norm_match(_as_text(set_row.get("longitudinal_label")))
                    and _norm_match(_as_text(row.get("base_longitudinal_label")))
                    == _norm_match(_as_text(span_default.get("base_label")))
                    and _norm_match(_as_text(row.get("additional_longitudinal_label")))
                    == _norm_match(_as_text(set_row.get("additional_label")))
                ),
                None,
            )
        if selected_row is None:
            selected_row = min(
                options,
                key=lambda row: (
                    _as_non_negative_float(row.get("weight_total_kg"))
                    if _as_non_negative_float(row.get("weight_total_kg")) is not None
                    else float("inf"),
                    _as_int(row.get("option")) if _as_int(row.get("option")) is not None else 999_999,
                ),
            )

        weight_total = float(_as_non_negative_float(selected_row.get("weight_total_kg")) or 0.0)
        span_total_weight += weight_total
        span_default["regions"].append(
            {
                "region_id": region_id,
                "option": _as_int(selected_row.get("option")),
                "transverse_label": _as_text(selected_row.get("transverse_label")),
                "longitudinal_label": _as_text(selected_row.get("longitudinal_label")),
                "base_longitudinal_label": _as_text(selected_row.get("base_longitudinal_label")) or "no se requiere",
                "additional_longitudinal_label": (
                    _as_text(selected_row.get("additional_longitudinal_label")) or "no se requiere"
                ),
                "weight_total_kg": weight_total,
                "weight_transverse_kg": float(_as_non_negative_float(selected_row.get("weight_transverse_kg")) or 0.0),
                "weight_longitudinal_kg": float(
                    _as_non_negative_float(selected_row.get("weight_longitudinal_kg")) or 0.0
                ),
            }
        )

    return span_default, span_total_weight


def _build_region_additional_options_by_base(
    region_row: dict[str, Any],
    base_options: list[dict[str, Any]],
    max_items: int = 10,
) -> dict[str, list[dict[str, Any]]]:
    options = region_row.get("options") if isinstance(region_row.get("options"), list) else []
    if not options or not base_options:
        return {}

    output: dict[str, list[dict[str, Any]]] = {}
    for base in base_options:
        base_value = _as_text(base.get("value"))
        base_label = _as_text(base.get("base_label")) or "no se requiere"
        rows = [
            row
            for row in options
            if (_as_text(row.get("base_longitudinal_label")) or "no se requiere") == base_label
        ]
        rows_sorted = sorted(
            rows,
            key=lambda item: (
                item.get("weight_additional_kg") if item.get("weight_additional_kg") is not None else float("inf"),
                item.get("option") if item.get("option") is not None else 999_999,
            ),
        )
        by_label: dict[str, dict[str, Any]] = {}
        for row in rows_sorted:
            add_label = _as_text(row.get("additional_longitudinal_label")) or "no se requiere"
            long_label = _as_text(row.get("longitudinal_label")) or add_label
            option_value = _as_int(row.get("option"))
            weight_kg = float(row.get("weight_additional_kg") or 0.0)
            base_region_weight_kg = float(row.get("weight_base_kg") or 0.0)
            long_total_weight_kg = _resolve_region_longitudinal_weight_kg(
                row,
                _as_float(region_row.get("length_mm")),
            )
            item = {
                "value": f"add_{option_value if option_value is not None else len(by_label) + 1}_{add_label.replace(' ', '_')}",
                "label": add_label,
                "longitudinal_label": long_label,
                "weight_kg": weight_kg,
                "weight_base_region_kg": base_region_weight_kg,
                "weight_longitudinal_total_kg": long_total_weight_kg,
                "option": option_value,
            }
            current = by_label.get(add_label)
            if current is None or (item["weight_kg"], item.get("option") or 999_999) < (
                current["weight_kg"],
                current.get("option") or 999_999,
            ):
                by_label[add_label] = item

        ordered = sorted(
            by_label.values(),
            key=lambda item: (item["weight_kg"], item.get("option") or 999_999, item["label"]),
        )
        output[base_value] = ordered[:max_items]
    return output


def _build_span_preview(
    span: dict[str, Any],
    *,
    optimized_map: dict[tuple[str, str], dict[str, Any]],
    schedule_map: dict[tuple[str, str], list[dict[str, Any]]],
    transverse_options_map: dict[tuple[str, str], list[dict[str, Any]]],
    default_span_length_mm: float = 6000.0,
) -> dict[str, Any]:
    span_id = _as_text(span.get("id")) or "S1"
    support_left_mm = _as_non_negative_float(span.get("support_left_mm"))
    support_right_mm = _as_non_negative_float(span.get("support_right_mm"))
    clear_length_mm = _as_float(span.get("clear_length_mm"))
    span_default_length_mm = clear_length_mm if clear_length_mm is not None else default_span_length_mm
    is_deep_beam = bool(span.get("is_deep_beam"))
    regions = span.get("regions") if isinstance(span.get("regions"), list) else []
    region_rows: list[dict[str, Any]] = []
    region_lengths: list[float] = []
    any_estimated = False
    span_width_mm: float | None = None
    span_height_mm: float | None = None
    span_d_mm: float | None = None

    for index, region in enumerate(regions, start=1):
        region_id = _as_text(region.get("id")) or f"R{index}"
        region_type = _as_text(region.get("type")) or "NC"
        start_ratio = _as_float(region.get("from")) or 0.0
        end_ratio = _as_float(region.get("to")) or 0.0
        if end_ratio < start_ratio:
            end_ratio = start_ratio
        ratio = max(end_ratio - start_ratio, 0.0)
        optimized = optimized_map.get((span_id, region_id))
        schedule_rows = schedule_map.get((span_id, region_id), [])
        sorted_schedule_rows = sorted(
            schedule_rows,
            key=lambda item: (
                item.get("weight_total_kg") if item.get("weight_total_kg") is not None else float("inf"),
                item.get("option") if item.get("option") is not None else 999_999,
            ),
        )
        schedule_row = _pick_schedule_row(schedule_rows, optimized)
        longitudinal_mode = _as_text((optimized or {}).get("longitudinal_mode")) or None
        schedule_option_limit = 200 if (longitudinal_mode or "").lower() == "span_coupled" else 10
        schedule_options = sorted_schedule_rows[:schedule_option_limit]
        transverse_options = list(transverse_options_map.get((span_id, region_id), []))
        if not transverse_options:
            transverse_options = _build_transverse_component_options(
                sorted_schedule_rows,
                max_items=10,
            )
        if (longitudinal_mode or "").lower() == "span_coupled":
            schedule_options = _expand_span_coupled_options(
                longitudinal_rows=schedule_options,
                transverse_options=transverse_options,
            )
            if schedule_options:
                if schedule_row is not None:
                    matched_selected = next(
                        (
                            row
                            for row in schedule_options
                            if _norm_match(_as_text(row.get("transverse_label")))
                            == _norm_match(_as_text(schedule_row.get("transverse_label")))
                            and _norm_match(_as_text(row.get("longitudinal_label")))
                            == _norm_match(_as_text(schedule_row.get("longitudinal_label")))
                        ),
                        None,
                    )
                    schedule_row = matched_selected or schedule_options[0]
                else:
                    schedule_row = schedule_options[0]
        longitudinal_options = _build_component_options(
            sorted_schedule_rows,
            label_key="longitudinal_label",
            weight_key="weight_longitudinal_kg",
        )

        if schedule_row and schedule_row.get("length_mm"):
            length_mm = float(schedule_row["length_mm"])
            length_estimated = False
        else:
            length_mm = max(span_default_length_mm * ratio, 1.0)
            length_estimated = clear_length_mm is None
            if length_estimated:
                any_estimated = True

        spacing_mm = optimized.get("spacing_mm") if optimized else None
        if spacing_mm is None:
            spacing_mm = 100 if region_type.upper() == "C" else 200
            spacing_estimated = True
        else:
            spacing_estimated = False

        transverse_label = (optimized or {}).get("transverse_label") or f"{region_type.upper()} @ {spacing_mm} mm est."
        longitudinal_label = (optimized or {}).get("longitudinal_label") or "long. n/d"
        if schedule_row is None and schedule_options:
            schedule_row = schedule_options[0]
        if schedule_row is not None:
            transverse_label = _as_text(schedule_row.get("transverse_label")) or transverse_label
            longitudinal_label = _as_text(schedule_row.get("longitudinal_label")) or longitudinal_label

        region_width_mm = _as_float(region.get("width_mm"))
        region_height_mm = _as_float(region.get("height_mm"))
        region_d_mm = _as_float(region.get("d_mm"))
        if region_width_mm is not None:
            span_width_mm = max(span_width_mm or 0.0, region_width_mm)
        if region_height_mm is not None:
            span_height_mm = max(span_height_mm or 0.0, region_height_mm)
        if region_d_mm is not None:
            span_d_mm = max(span_d_mm or 0.0, region_d_mm)

        region_rows.append(
            {
                "region_id": region_id,
                "type": region_type.upper(),
                "from": start_ratio,
                "to": end_ratio,
                "length_mm": int(round(length_mm)),
                "length_estimated": length_estimated,
                "spacing_mm": int(round(float(spacing_mm))),
                "spacing_estimated": spacing_estimated,
                "transverse_label": transverse_label,
                "longitudinal_label": longitudinal_label,
                "base_longitudinal_label": (
                    _as_text((schedule_row or {}).get("base_longitudinal_label"))
                    or _long_arrangement_label(
                        _as_text((optimized or {}).get("base_long_bar")),
                        _as_int((optimized or {}).get("base_long_count")),
                        empty_label="no se requiere",
                    )
                    or None
                ),
                "additional_longitudinal_label": (
                    _as_text((schedule_row or {}).get("additional_longitudinal_label"))
                    or _long_arrangement_label(
                        _as_text((optimized or {}).get("extra_long_bar")),
                        _as_int((optimized or {}).get("extra_long_count")),
                        empty_label="no se requiere",
                    )
                ),
                "base_long_bar": _as_text((schedule_row or {}).get("base_long_bar"))
                or _as_text((optimized or {}).get("base_long_bar"))
                or None,
                "base_long_count": _as_int((schedule_row or {}).get("base_long_count"))
                or _as_int((optimized or {}).get("base_long_count")),
                "extra_long_bar": _as_text((schedule_row or {}).get("extra_long_bar"))
                or _as_text((optimized or {}).get("extra_long_bar"))
                or None,
                "extra_long_count": _as_int((schedule_row or {}).get("extra_long_count"))
                or _as_int((optimized or {}).get("extra_long_count")),
                "longitudinal_mode": longitudinal_mode,
                "is_deep_beam": is_deep_beam,
                "selected_option": _as_int(schedule_row.get("option")) if schedule_row is not None else None,
                "best_option": _as_int(schedule_options[0].get("option")) if schedule_options else None,
                "best_weight_kg": (
                    float(schedule_options[0]["weight_total_kg"])
                    if schedule_options and schedule_options[0].get("weight_total_kg") is not None
                    else None
                ),
                "selected_weight_kg": (
                    float(schedule_row["weight_total_kg"])
                    if schedule_row is not None and schedule_row.get("weight_total_kg") is not None
                    else None
                ),
                "options": [
                    {
                        "option": _as_int(opt.get("option")) or (opt_index + 1),
                        "transverse_label": _as_text(opt.get("transverse_label")),
                        "longitudinal_label": _as_text(opt.get("longitudinal_label")),
                        "base_longitudinal_label": _as_text(opt.get("base_longitudinal_label"))
                        or _long_arrangement_label(
                            _as_text(opt.get("base_long_bar")),
                            _as_int(opt.get("base_long_count")),
                            empty_label="no se requiere",
                        ),
                        "additional_longitudinal_label": _as_text(opt.get("additional_longitudinal_label"))
                        or _long_arrangement_label(
                            _as_text(opt.get("extra_long_bar")),
                            _as_int(opt.get("extra_long_count")),
                            empty_label="no se requiere",
                        ),
                        "base_long_bar": _as_text(opt.get("base_long_bar")) or None,
                        "base_long_count": _as_int(opt.get("base_long_count")),
                        "extra_long_bar": _as_text(opt.get("extra_long_bar")) or None,
                        "extra_long_count": _as_int(opt.get("extra_long_count")),
                        "stirrup_count": _as_non_negative_int(opt.get("stirrup_count")),
                        "stirrup_unit_weight_kg": (
                            float(opt["stirrup_unit_weight_kg"])
                            if opt.get("stirrup_unit_weight_kg") is not None
                            else None
                        ),
                        "weight_transverse_kg": (
                            float(opt["weight_transverse_kg"])
                            if opt.get("weight_transverse_kg") is not None
                            else None
                        ),
                        "weight_longitudinal_kg": (
                            float(opt["weight_longitudinal_kg"])
                            if opt.get("weight_longitudinal_kg") is not None
                            else None
                        ),
                        "weight_base_kg": (
                            float(opt["weight_base_kg"])
                            if opt.get("weight_base_kg") is not None
                            else None
                        ),
                        "weight_additional_kg": (
                            float(opt["weight_additional_kg"])
                            if opt.get("weight_additional_kg") is not None
                            else None
                        ),
                        "weight_total_kg": (
                            float(opt["weight_total_kg"])
                            if opt.get("weight_total_kg") is not None
                            else None
                        ),
                    }
                    for opt_index, opt in enumerate(schedule_options)
                ],
                "transverse_options": transverse_options,
                "longitudinal_options": longitudinal_options,
            }
        )
        region_lengths.append(length_mm)

    longitudinal_base_options = _build_span_longitudinal_base_options(region_rows)
    for row in region_rows:
        row["additional_options_by_base"] = _build_region_additional_options_by_base(
            row,
            longitudinal_base_options,
        )
    span_longitudinal_option_sets = _build_span_longitudinal_option_sets(
        region_rows,
        longitudinal_base_options,
    )
    default_longitudinal_set_value: str | None = None
    default_longitudinal_base_value: str | None = None
    if span_longitudinal_option_sets:
        best_set_by_total = min(
            span_longitudinal_option_sets,
            key=lambda item: (
                _as_non_negative_float(item.get("total_weight_kg"))
                if _as_non_negative_float(item.get("total_weight_kg")) is not None
                else float("inf"),
                _as_non_negative_float(item.get("total_longitudinal_weight_kg"))
                if _as_non_negative_float(item.get("total_longitudinal_weight_kg")) is not None
                else float("inf"),
            ),
        )
        default_longitudinal_set_value = _as_text(best_set_by_total.get("value")) or None
        default_longitudinal_base_value = _as_text(best_set_by_total.get("base_value")) or None
    if not default_longitudinal_base_value:
        default_longitudinal_base_value = (
            _as_text(longitudinal_base_options[0].get("value")) if longitudinal_base_options else None
        )

    span_length_mm = int(round(sum(region_lengths))) if region_lengths else int(span_default_length_mm)
    if not region_rows and clear_length_mm is None:
        any_estimated = True
    return {
        "span_id": span_id,
        "seismic": _as_text(span.get("seismic")),
        "gravity": _as_text(span.get("gravity")),
        "support_left_mm": int(round(support_left_mm)) if support_left_mm is not None else None,
        "support_right_mm": int(round(support_right_mm)) if support_right_mm is not None else None,
        "length_mm": span_length_mm,
        "length_estimated": any_estimated,
        "clear_length_mm": int(round(clear_length_mm)) if clear_length_mm is not None else None,
        "width_mm": int(round(span_width_mm)) if span_width_mm is not None else None,
        "height_mm": int(round(span_height_mm)) if span_height_mm is not None else None,
        "d_mm": int(round(span_d_mm)) if span_d_mm is not None else None,
        "is_deep_beam": is_deep_beam,
        "longitudinal_base_options": longitudinal_base_options,
        "default_longitudinal_base_value": default_longitudinal_base_value,
        "span_longitudinal_option_sets": span_longitudinal_option_sets,
        "default_longitudinal_option_set_value": default_longitudinal_set_value,
        "span_option_choices": _build_span_option_choices(region_rows),
        "regions": region_rows,
    }


def build_job_preview_payload(job_id: str) -> dict[str, Any]:
    meta = get_job(job_id)
    case_payload = get_job_case_payload(job_id)
    artifacts = meta.get("artifacts", {})
    optimized_path = Path(artifacts["optimized_results.xlsx"]) if "optimized_results.xlsx" in artifacts else None
    schedule_path = Path(artifacts["reinforcement_schedule.xlsx"]) if "reinforcement_schedule.xlsx" in artifacts else None

    cache_key = (
        str(meta.get("status") or ""),
        _artifact_mtime(optimized_path),
        _artifact_mtime(schedule_path),
    )
    with _PREVIEW_CACHE_LOCK:
        cached = _PREVIEW_CACHE.get(job_id)
    if cached is not None and cached[0] == cache_key:
        return cached[1]

    optimized_map = _read_optimized_regions(optimized_path)
    schedule_map = _read_schedule_rows(schedule_path)
    transverse_options_map = _read_transverse_options(schedule_path)

    beams = case_payload.get("beams") if isinstance(case_payload.get("beams"), list) else []
    beam = beams[0] if beams else {}
    spans = beam.get("spans") if isinstance(beam.get("spans"), list) else []
    span_rows = [
        _build_span_preview(
            span,
            optimized_map=optimized_map,
            schedule_map=schedule_map,
            transverse_options_map=transverse_options_map,
        )
        for span in spans
    ]
    default_spans: list[dict[str, Any]] = []
    optimal_beam_weight_kg = 0.0
    for span_row in span_rows:
        span_default, span_weight = _resolve_span_default_selection(span_row)
        default_spans.append(span_default)
        optimal_beam_weight_kg += float(span_weight or 0.0)

    total_span_mm = sum(span_row.get("length_mm", 0) for span_row in span_rows)
    total_support_mm = sum(
        int(span_row.get("support_right_mm", 0) or 0)
        for span_row in span_rows
    )
    if total_span_mm <= 0:
        total_span_mm = 6000

    beam_width_from_spans = max(
        (int(span_row.get("width_mm", 0) or 0) for span_row in span_rows),
        default=0,
    )
    beam_height_from_spans = max(
        (int(span_row.get("height_mm", 0) or 0) for span_row in span_rows),
        default=0,
    )

    payload = {
        "job_id": job_id,
        "status": meta.get("status"),
        "beam": {
            "beam_id": beam.get("beam_id"),
            "detailing": beam.get("detailing"),
            "fc_mpa": beam.get("fc_mpa"),
            "fy_mpa": beam.get("fy_mpa"),
            "width_mm": beam.get("width_mm") or (beam_width_from_spans or None),
            "height_mm": beam.get("height_mm") or (beam_height_from_spans or None),
        },
        "spans": span_rows,
        "optimal_beam_weight_kg": float(optimal_beam_weight_kg),
        "default_selection": {
            "spans": default_spans,
            "total_weight_kg": float(optimal_beam_weight_kg),
        },
        "total_span_mm": int(round(total_span_mm)),
        "total_support_mm": int(round(total_support_mm)),
        "total_system_mm": int(round(total_span_mm + total_support_mm)),
        "total_span_estimated": any(span_row.get("length_estimated", True) for span_row in span_rows),
        "data_source": {
            "optimized_results": bool(optimized_map),
            "reinforcement_schedule": bool(schedule_map),
        },
    }
    with _PREVIEW_CACHE_LOCK:
        _PREVIEW_CACHE[job_id] = (cache_key, payload)
    return payload
