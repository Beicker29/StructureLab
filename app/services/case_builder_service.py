from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from fastapi import UploadFile
from openpyxl import load_workbook

from app.core.errors import DomainValidationAppError, InvalidUploadError
from app.services.case_service import (
    ensure_upload_suffix,
    merge_optimization_defaults,
    read_upload_bytes,
)
from rc_shear_torsion.domain.errors import DomainValidationError
from rc_shear_torsion.domain.validation import validate_case_payload
from rc_shear_torsion.io import REQUIRED_COLUMNS, as_text, find_header


def _natural_name_key(value: str) -> tuple[int, Any]:
    text = value.strip()
    if text.isdigit():
        return (0, int(text))
    return (1, text)


def _extract_unique_names(excel_bytes: bytes, sheet_name: str, field_name: str) -> set[str]:
    try:
        workbook = load_workbook(filename=BytesIO(excel_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidUploadError(f"No fue posible leer '{field_name}' como archivo Excel valido") from exc

    if sheet_name not in workbook.sheetnames:
        raise InvalidUploadError(f"La hoja '{sheet_name}' no existe en '{field_name}'")

    worksheet = workbook[sheet_name]
    try:
        header_row_idx, header_map = find_header(worksheet.iter_rows(min_row=1, max_row=30, values_only=True))
    except ValueError as exc:
        raise InvalidUploadError(
            f"'{field_name}' no tiene encabezado ETABS valido en las primeras 30 filas"
        ) from exc
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing_columns:
        raise InvalidUploadError(
            f"'{field_name}' no contiene columnas requeridas: {', '.join(missing_columns)}"
        )

    names: set[str] = set()
    for row in worksheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None:
            continue
        unique_name = as_text(row[header_map["UniqueName"]])
        station_value = row[header_map["Station"]]
        if not unique_name or station_value in (None, ""):
            continue
        names.add(unique_name)

    if not names:
        raise InvalidUploadError(
            f"'{field_name}' no contiene filas validas con columnas UniqueName y Station"
        )
    return names


def _parse_csv_names(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [item.strip() for item in raw.split(",") if item.strip()]


def _parse_pairs_json(raw: str | None) -> list[dict[str, str]]:
    if not raw:
        return []
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise InvalidUploadError("frame_pairs_json debe ser un JSON valido") from exc

    if not isinstance(payload, list):
        raise InvalidUploadError("frame_pairs_json debe ser una lista de objetos")

    parsed_pairs: list[dict[str, str]] = []
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise InvalidUploadError(f"frame_pairs_json[{index}] debe ser un objeto")
        seismic = str(item.get("seismic", "")).strip()
        gravity = str(item.get("gravity", "")).strip()
        span_id = str(item.get("id", "")).strip() or f"S{index}"
        if not seismic or not gravity:
            raise InvalidUploadError(
                f"frame_pairs_json[{index}] requiere 'seismic' y 'gravity'"
            )
        parsed_pairs.append({"id": span_id, "seismic": seismic, "gravity": gravity})
    return parsed_pairs


def _parse_optimization_overrides(raw: str | None) -> dict[str, Any] | None:
    if not raw:
        return None
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise InvalidUploadError("optimization_overrides_json debe ser JSON valido") from exc
    if not isinstance(payload, dict):
        raise InvalidUploadError("optimization_overrides_json debe ser un objeto JSON")
    return payload


def build_case_payload_from_form(
    *,
    seismic_excel: UploadFile,
    gravity_excel: UploadFile,
    max_upload_bytes: int,
    case_name: str,
    sheet_name: str,
    units_rebar_per_length: str,
    beam_id: str,
    detailing: str,
    cover_side_mm: float,
    cover_top_mm: float,
    cover_bottom_mm: float,
    fc_mpa: float,
    fy_mpa: float,
    width_mm: float,
    height_mm: float,
    d_mm: float,
    db_bar: str,
    min_branches_c: int,
    min_branches_nc: int,
    region_c_ratio: float,
    frame_names_csv: str | None,
    frame_pairs_json: str | None,
    optimization_overrides_json: str | None,
) -> dict[str, Any]:
    case_name = case_name.strip() or "case_from_form"
    sheet_name = sheet_name.strip()
    if not sheet_name:
        raise InvalidUploadError("sheet_name no puede estar vacio")
    beam_id = beam_id.strip() or "B1"

    if region_c_ratio <= 0.0 or region_c_ratio >= 0.5:
        raise InvalidUploadError("region_c_ratio debe estar entre 0 y 0.5")

    ensure_upload_suffix(seismic_excel, {".xlsx"}, "seismic_excel")
    ensure_upload_suffix(gravity_excel, {".xlsx"}, "gravity_excel")
    seismic_bytes = read_upload_bytes(seismic_excel, max_upload_bytes)
    gravity_bytes = read_upload_bytes(gravity_excel, max_upload_bytes)

    seismic_names = _extract_unique_names(seismic_bytes, sheet_name, "seismic_excel")
    gravity_names = _extract_unique_names(gravity_bytes, sheet_name, "gravity_excel")

    span_pairs = _parse_pairs_json(frame_pairs_json)
    if span_pairs:
        for pair in span_pairs:
            if pair["seismic"] not in seismic_names:
                raise InvalidUploadError(f"El vano seismic '{pair['seismic']}' no existe en seismic_excel")
            if pair["gravity"] not in gravity_names:
                raise InvalidUploadError(f"El vano gravity '{pair['gravity']}' no existe en gravity_excel")
    else:
        csv_names = _parse_csv_names(frame_names_csv)
        if csv_names:
            for frame_name in csv_names:
                if frame_name not in seismic_names:
                    raise InvalidUploadError(f"El vano '{frame_name}' no existe en seismic_excel")
                if frame_name not in gravity_names:
                    raise InvalidUploadError(f"El vano '{frame_name}' no existe en gravity_excel")
            span_pairs = [
                {"id": f"S{index}", "seismic": frame_name, "gravity": frame_name}
                for index, frame_name in enumerate(csv_names, start=1)
            ]
        else:
            shared = sorted(seismic_names & gravity_names, key=_natural_name_key)
            if not shared:
                raise InvalidUploadError(
                    "No se encontraron UniqueName comunes entre seismic_excel y gravity_excel. "
                    "Usa frame_pairs_json para mapear vanos distintos."
                )
            span_pairs = [
                {"id": f"S{index}", "seismic": frame_name, "gravity": frame_name}
                for index, frame_name in enumerate(shared, start=1)
            ]

    optimization_overrides = _parse_optimization_overrides(optimization_overrides_json)
    optimization_payload = merge_optimization_defaults(optimization_overrides)

    spans: list[dict[str, Any]] = []
    middle_to = 1.0 - region_c_ratio
    for index, pair in enumerate(span_pairs, start=1):
        span_id = f"{beam_id}.{index}"
        spans.append(
            {
                "id": span_id,
                "seismic": pair["seismic"],
                "gravity": pair["gravity"],
                "regions": [
                    {
                        "id": "R1",
                        "from": 0.0,
                        "to": region_c_ratio,
                        "type": "C",
                        "d_mm": d_mm,
                        "db_bar": db_bar,
                        "min_branches": min_branches_c,
                        "width_mm": width_mm,
                        "height_mm": height_mm,
                    },
                    {
                        "id": "R2",
                        "from": region_c_ratio,
                        "to": middle_to,
                        "type": "NC",
                        "d_mm": d_mm,
                        "db_bar": db_bar,
                        "min_branches": min_branches_nc,
                        "width_mm": width_mm,
                        "height_mm": height_mm,
                    },
                    {
                        "id": "R3",
                        "from": middle_to,
                        "to": 1.0,
                        "type": "C",
                        "d_mm": d_mm,
                        "db_bar": db_bar,
                        "min_branches": min_branches_c,
                        "width_mm": width_mm,
                        "height_mm": height_mm,
                    },
                ],
            }
        )

    case_payload = {
        "case_name": case_name,
        "inputs": {
            "seismic_excel": "seismic.xlsx",
            "gravity_excel": "gravity.xlsx",
            "sheet_name": sheet_name,
        },
        "units": {"rebar_per_length": units_rebar_per_length},
        "beams": [
            {
                "beam_id": beam_id,
                "detailing": detailing,
                "cover_side_mm": cover_side_mm,
                "cover_top_mm": cover_top_mm,
                "cover_bottom_mm": cover_bottom_mm,
                "fc_mpa": fc_mpa,
                "fy_mpa": fy_mpa,
                "spans": spans,
            }
        ],
        "optimization": optimization_payload,
    }
    try:
        validated = validate_case_payload(case_payload)
    except DomainValidationError as exc:
        raise DomainValidationAppError(
            message="El formulario no cumple reglas de negocio/ingenieria",
            details=exc.to_dicts(),
        ) from exc

    return validated.model_dump(mode="python", by_alias=True)
