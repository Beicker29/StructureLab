from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.services.job_preview_service import _build_span_preview, _read_transverse_options
from rc_shear_torsion.design import (
    DemandScenario,
    RegionDemand,
    classify_torsion_state,
    optimize_region_exhaustive,
    optimize_region_ga,
    top_region_alternatives,
)
from rc_shear_torsion.models import OptimizationConfig, VariablesConfig
from rc_shear_torsion.ranking import transverse_alternative_rank_key
from rc_shear_torsion.report import write_reinforcement_schedule


ROOT = Path(__file__).resolve().parents[1]


def _region() -> RegionDemand:
    scenario = DemandScenario(
        source="SEISMIC",
        station_mm=0.0,
        x_relative=0.5,
        region_id="R1",
        v_rebar_mm2_per_m=1.0,
        t_transverse_mm2_per_m=0.0,
        t_longitudinal_mm2=0.0,
        torsion_state=classify_torsion_state(0.0, 0.0),
    )
    return RegionDemand(
        beam_id="B1",
        span_id="S1",
        region_id="R1",
        region_type="NC",
        beam_detailing="DMI",
        d_mm=600.0,
        db_bar=None,
        min_branches=2,
        width_mm=300.0,
        height_mm=600.0,
        cover_side_mm=40.0,
        cover_top_mm=40.0,
        cover_bottom_mm=40.0,
        scenarios=(scenario,),
        fc_mpa=28.0,
        fy_mpa=420.0,
        region_length_mm=400.0,
    )


def _variables() -> VariablesConfig:
    return VariablesConfig(
        E_bars=["#3"],
        G_bars=["#3"],
        G_counts=[0],
        stirrup_spacing_mm=[100, 110, 120, 130, 140],
        stirrup_spacing_project_max_mm=130,
        longitudinal_bars=["#4"],
        longitudinal_bar_counts=[2],
    )


def _ga_config(variables: VariablesConfig) -> OptimizationConfig:
    return OptimizationConfig.model_validate(
        {
            "enabled": True,
            "objective": "min_weight",
            "variables": variables.model_dump(),
            "genetic_algorithm": {
                "population_size": 60,
                "generations": 20,
                "crossover_rate": 0.8,
                "mutation_rate": 0.3,
                "elite_count": 5,
            },
        }
    )


class TransverseRankingTests(unittest.TestCase):
    def test_same_weight_prefers_larger_spacing_before_heavier_candidate(self) -> None:
        candidates = [
            {"id": "A", "spacing": 110, "weight": 3.03},
            {"id": "B", "spacing": 120, "weight": 3.03},
            {"id": "C", "spacing": 130, "weight": 3.03},
            {"id": "D", "spacing": 100, "weight": 3.79},
        ]

        ordered = sorted(
            candidates,
            key=lambda item: transverse_alternative_rank_key(
                weight_kg=item["weight"],
                spacing_mm=item["spacing"],
                stable_tie_break=(item["id"],),
            ),
        )

        self.assertEqual([item["id"] for item in ordered], ["C", "B", "A", "D"])

    def test_lower_weight_has_priority_over_larger_spacing(self) -> None:
        candidates = [
            {"id": "A", "spacing": 150, "weight": 4.0},
            {"id": "B", "spacing": 100, "weight": 3.0},
        ]

        ordered = sorted(
            candidates,
            key=lambda item: transverse_alternative_rank_key(
                weight_kg=item["weight"],
                spacing_mm=item["spacing"],
                stable_tie_break=(item["id"],),
            ),
        )

        self.assertEqual([item["id"] for item in ordered], ["B", "A"])

    def test_exact_tie_uses_only_stable_deterministic_criteria(self) -> None:
        candidates = ["stable-b", "stable-a"]

        ordered = sorted(
            candidates,
            key=lambda stable_id: transverse_alternative_rank_key(
                weight_kg=3.03,
                spacing_mm=130,
                stable_tie_break=(stable_id,),
            ),
        )

        self.assertEqual(ordered, ["stable-a", "stable-b"])

    def test_catalog_omits_infeasible_candidate_and_orders_equal_weight_by_spacing(self) -> None:
        alternatives = top_region_alternatives(
            _region(),
            _variables(),
            top_n=None,
            check_longitudinal=False,
        )

        self.assertEqual([item.spacing_mm for item in alternatives], [130, 120, 110, 100])
        self.assertNotIn(140, [item.spacing_mm for item in alternatives])
        self.assertTrue(all(item.status == "ok" for item in alternatives))

    def test_ga_and_exhaustive_choose_same_ranked_candidate(self) -> None:
        region = _region()
        variables = _variables()

        exhaustive = optimize_region_exhaustive(region, variables, check_longitudinal=False)
        genetic = optimize_region_ga(region, _ga_config(variables), check_longitudinal=False)

        self.assertEqual(exhaustive.selected.spacing_mm, 130)
        self.assertEqual(genetic.selected.spacing_mm, exhaustive.selected.spacing_mm)
        self.assertEqual(genetic.selected.objective, exhaustive.selected.objective)

    def test_report_and_preview_keep_optimal_transverse_candidate_first(self) -> None:
        region = _region()
        variables = _variables()
        alternatives = top_region_alternatives(
            region,
            variables,
            top_n=None,
            check_longitudinal=False,
        )
        optimal = optimize_region_exhaustive(region, variables, check_longitudinal=False).selected
        runtime_root = ROOT / ".tmp_test_runtime"
        runtime_root.mkdir(exist_ok=True)

        with tempfile.TemporaryDirectory(dir=runtime_root) as temp_dir:
            schedule_path = Path(temp_dir) / "reinforcement_schedule.xlsx"
            key = ("B1", "S1", "R1")
            write_reinforcement_schedule(
                schedule_path,
                [alternatives[0]],
                region_lengths_mm={key: 401.0},
                transverse_alternatives={key: alternatives},
            )

            workbook = load_workbook(schedule_path, data_only=True, read_only=True)
            try:
                rows = list(workbook["transversales"].iter_rows(values_only=True))
            finally:
                workbook.close()
            header = {str(value): index for index, value in enumerate(rows[0])}
            reported_spacings = [int(row[header["espaciamiento_mm"]]) for row in rows[1:]]
            self.assertEqual(reported_spacings, [130, 120, 110, 100])

            transverse_map = _read_transverse_options(schedule_path)
            preview = _build_span_preview(
                {
                    "id": "S1",
                    "clear_length_mm": 401.0,
                    "regions": [{"id": "R1", "type": "NC", "from": 0.0, "to": 1.0}],
                },
                optimized_map={
                    ("S1", "R1"): {
                        "spacing_mm": optimal.spacing_mm,
                        "transverse_label": f"1E #3 @ {optimal.spacing_mm} mm",
                        "longitudinal_label": "no se requiere",
                    }
                },
                schedule_map={},
                transverse_options_map=transverse_map,
            )

        first_region = preview["regions"][0]
        self.assertEqual(int(first_region["transverse_options"][0]["spacing_mm"]), optimal.spacing_mm)
        self.assertEqual(first_region["transverse_label"], first_region["transverse_options"][0]["label"])

    def test_ui_preserves_backend_order_and_existing_valid_manual_selection(self) -> None:
        javascript = (ROOT / "app" / "ui" / "static" / "ui.js").read_text(encoding="utf-8")

        self.assertIn(
            "const transOptionsSource=(Array.isArray(rg.transverse_options)?rg.transverse_options:[])"
            ".map(",
            javascript,
        )
        self.assertNotIn("filter(opt=>opt.label).sort((a,b)=>{const aw=", javascript)
        self.assertIn(
            "const initialTrans=region.transOptions.some(opt=>opt.label===saved.transverse_label)"
            "?saved.transverse_label:region.defaultTrans;",
            javascript,
        )


if __name__ == "__main__":
    unittest.main()
