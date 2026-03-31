from __future__ import annotations

import uuid
import unittest
from pathlib import Path

from openpyxl import load_workbook

from rc_shear_torsion.design import RegionDesignResult
from rc_shear_torsion.report import _stirrup_count_for_reporting, write_reinforcement_schedule


class ReportTests(unittest.TestCase):
    def _region_result(self, *, span_id: str, region_id: str, spacing_mm: int = 100) -> RegionDesignResult:
        return RegionDesignResult(
            beam_id="B1",
            span_id=span_id,
            region_id=region_id,
            region_type="C",
            source_control="seismic",
            governing_station=0.0,
            v_req=100.0,
            t_req=50.0,
            l_req=0.0,
            e_bar="#3",
            g_bar="#3",
            g_count=2,
            spacing_mm=spacing_mm,
            av1=100.0,
            av2=142.0,
            av_total=242.0,
            at=71.0,
            at_over_s=710.0,
            av_over_s=2420.0,
            long_bar="",
            long_count=0,
            controlling_limit="Resistencia",
            failure_mode="ok",
            status="ok",
            message="ok",
            objective=0.0,
            method="test",
            evaluated_candidates=1,
            feasible_candidates=1,
            transverse_weight_kg_per_m=1.0,
            longitudinal_weight_kg_per_m=0.0,
            stirrup_unit_weight_kg=0.5,
        )

    def test_stirrup_count_support_offsets(self) -> None:
        self.assertEqual(
            _stirrup_count_for_reporting(6000.0, 150, is_first_region=True, is_last_region=True),
            41,
        )
        self.assertEqual(
            _stirrup_count_for_reporting(1200.0, 100, is_first_region=True, is_last_region=False),
            12,
        )
        self.assertEqual(
            _stirrup_count_for_reporting(1200.0, 100, is_first_region=False, is_last_region=False),
            13,
        )
        self.assertEqual(
            _stirrup_count_for_reporting(1200.0, 100, is_first_region=False, is_last_region=True),
            12,
        )

    def test_reinforcement_schedule_uses_offsets_in_span_edges(self) -> None:
        results = [
            self._region_result(span_id="S1", region_id="R1", spacing_mm=100),
            self._region_result(span_id="S1", region_id="R2", spacing_mm=100),
            self._region_result(span_id="S1", region_id="R3", spacing_mm=100),
        ]
        lengths = {
            ("B1", "S1", "R1"): 1000.0,
            ("B1", "S1", "R2"): 1000.0,
            ("B1", "S1", "R3"): 1000.0,
        }

        tmp_root = Path(__file__).resolve().parents[1] / ".tmp_test_report"
        tmp_root.mkdir(parents=True, exist_ok=True)
        out_path = tmp_root / f"reinforcement_schedule_{uuid.uuid4().hex}.xlsx"
        try:
            write_reinforcement_schedule(out_path, results, region_lengths_mm=lengths)

            workbook = load_workbook(out_path, data_only=True, read_only=True)
            try:
                sheet = workbook["por_region"]
                rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()

            header = {str(name): idx for idx, name in enumerate(rows[0])}
            region_idx = header["region_id"]
            option_idx = header["opcion"]
            count_idx = header["cantidad_estribos_region"]

            counts_by_region: dict[str, int] = {}
            for row in rows[1:]:
                if int(row[option_idx]) != 1:
                    continue
                counts_by_region[str(row[region_idx])] = int(row[count_idx])

            self.assertEqual(counts_by_region["R1"], 10)
            self.assertEqual(counts_by_region["R2"], 11)
            self.assertEqual(counts_by_region["R3"], 10)
        finally:
            if out_path.exists():
                out_path.unlink()


if __name__ == "__main__":
    unittest.main()

