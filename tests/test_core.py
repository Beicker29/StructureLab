from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import unittest
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from rc_shear_torsion.domain.errors import DomainValidationError
from rc_shear_torsion.domain.validation import validate_case_payload
from rc_shear_torsion.design import RegionDemand, evaluate_candidate, optimize_region_exhaustive
from rc_shear_torsion.models import CaseConfig, VariablesConfig, resolve_path
from rc_shear_torsion.run import run_case

TMP_TEST_ROOT = Path(__file__).resolve().parents[1] / ".tmp_test_runtime"
TMP_TEST_ROOT.mkdir(parents=True, exist_ok=True)


def _new_runtime_dir(prefix: str) -> Path:
    path = TMP_TEST_ROOT / f"{prefix}_{uuid4().hex}"
    path.mkdir(parents=True, exist_ok=False)
    return path


class CoreTests(unittest.TestCase):
    def test_region_validation_rejects_gaps(self) -> None:
        payload = {
            "case_name": "bad_case",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {"id": "R1", "from": 0.0, "to": 0.3, "type": "C"},
                                {"id": "R2", "from": 0.4, "to": 1.0, "type": "NC"},
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        with self.assertRaises(ValueError):
            CaseConfig.model_validate(payload)

    def test_candidate_formula(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="NC",
            beam_detailing="DES",
            d_mm=None,
            db_bar=None,
            min_branches=None,
            width_mm=None,
            height_mm=None,
            cover_side_mm=None,
            cover_top_mm=None,
            cover_bottom_mm=None,
            source_control="seismic",
            governing_station=1.0,
            v_req=500.0,
            t_req=300.0,
            l_req=400.0,
            station_count=5,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#4",
            g_bar="#3",
            g_count=2,
            spacing_mm=100,
            long_bar="#5",
            long_count=4,
        )
        self.assertEqual(candidate.status, "ok")
        self.assertGreaterEqual(candidate.at_over_s, region.t_req)
        self.assertGreaterEqual(candidate.av_over_s, region.v_req)
        self.assertGreaterEqual(candidate.long_provided, region.l_req)

    def test_dmo_requires_detail_inputs_in_confined_region(self) -> None:
        payload = {
            "case_name": "dmo_missing_inputs",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DMO",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                    "d_mm": 600.0,
                                    "db_bar": "#5"
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        with self.assertRaises(DomainValidationError):
            validate_case_payload(payload)

    def test_des_requires_detail_inputs_in_confined_region(self) -> None:
        payload = {
            "case_name": "des_missing_inputs",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DES",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        with self.assertRaises(DomainValidationError):
            validate_case_payload(payload)

    def test_dmo_non_confined_requires_fc_fy_and_detail_inputs(self) -> None:
        payload = {
            "case_name": "dmo_nc_missing_fc_fy",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DMO",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "NC",
                                    "d_mm": 600.0,
                                    "db_bar": "#5",
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        with self.assertRaises(DomainValidationError):
            validate_case_payload(payload)

    def test_bar_range_allows_from_2_to_11(self) -> None:
        payload = {
            "case_name": "bar_range_ok",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DES",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                    "d_mm": 600.0,
                                    "db_bar": "#11",
                                    "width_mm": 300.0,
                                    "height_mm": 600.0,
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#2", "#11"],
                    "G_bars": ["#2", "#11"],
                    "G_counts": [0, 1],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#2", "#11"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        cfg = validate_case_payload(payload)
        self.assertEqual(cfg.beams[0].spans[0].regions[0].db_bar, "#11")

    def test_dmo_confined_region_detailing_rule(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=3,
            width_mm=None,
            height_mm=None,
            cover_side_mm=None,
            cover_top_mm=None,
            cover_bottom_mm=None,
            source_control="seismic",
            governing_station=1.0,
            v_req=50.0,
            t_req=20.0,
            l_req=200.0,
            station_count=5,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#3",
            g_bar="#3",
            g_count=0,
            spacing_mm=120,
            long_bar="#5",
            long_count=8,
        )
        self.assertEqual(candidate.status, "fail")
        self.assertEqual(candidate.failure_mode, "region_detail_fail")
        self.assertIn("minimum branches", candidate.message)

    def test_dmo_region_torsion_requires_at_least_two_user_branches(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=1,
            width_mm=300.0,
            height_mm=600.0,
            cover_side_mm=40.0,
            cover_top_mm=40.0,
            cover_bottom_mm=40.0,
            source_control="gravity",
            governing_station=0.0,
            v_req=10.0,
            t_req=100.0,
            l_req=10.0,
            station_count=3,
            fc_mpa=28.0,
            fy_mpa=420.0,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#4",
            g_bar="#3",
            g_count=0,
            spacing_mm=100,
            long_bar="#4",
            long_count=2,
        )
        self.assertEqual(candidate.status, "fail")
        self.assertEqual(candidate.failure_mode, "region_detail_fail")
        self.assertIn("TTrnRebar>0", candidate.message)
        self.assertIn("min_branches >= 2", candidate.message)

    def test_min_branches_is_enforced_in_optimization_domain(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=5,
            width_mm=None,
            height_mm=None,
            cover_side_mm=None,
            cover_top_mm=None,
            cover_bottom_mm=None,
            source_control="gravity",
            governing_station=0.0,
            v_req=3000.0,
            t_req=200.0,
            l_req=100.0,
            station_count=3,
        )
        variables = VariablesConfig(
            E_bars=["#4"],
            G_bars=["#3"],
            G_counts=[0, 1, 2, 3, 4, 5],
            stirrup_spacing_mm=[150],
            longitudinal_bars=["#4"],
            longitudinal_bar_counts=[2],
        )
        outcome = optimize_region_exhaustive(region, variables)
        self.assertGreaterEqual(outcome.selected.g_count, 3)

    def test_min_branches_without_valid_g_count_returns_input_fail(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=8,
            width_mm=None,
            height_mm=None,
            cover_side_mm=None,
            cover_top_mm=None,
            cover_bottom_mm=None,
            source_control="gravity",
            governing_station=0.0,
            v_req=10.0,
            t_req=10.0,
            l_req=10.0,
            station_count=3,
        )
        variables = VariablesConfig(
            E_bars=["#4"],
            G_bars=["#3"],
            G_counts=[0, 1, 2, 3, 4, 5],
            stirrup_spacing_mm=[150],
            longitudinal_bars=["#4"],
            longitudinal_bar_counts=[2],
        )
        outcome = optimize_region_exhaustive(region, variables)
        self.assertEqual(outcome.selected.failure_mode, "input_fail")
        self.assertEqual(outcome.evaluated_candidates, 0)
        self.assertIn("No G_counts satisfy min_branches", outcome.selected.message)

    def test_dmo_spacing_limit_reports_24dest_term(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DMO",
            d_mm=1000.0,
            db_bar="#5",
            min_branches=4,
            width_mm=None,
            height_mm=None,
            cover_side_mm=None,
            cover_top_mm=None,
            cover_bottom_mm=None,
            source_control="gravity",
            governing_station=0.0,
            v_req=10.0,
            t_req=10.0,
            l_req=10.0,
            station_count=3,
        )
        # #3 -> dest=9.5 mm; 24*dest=228 mm. En este dominio de barras, 150 mm suele controlar.
        # Esta prueba valida que 24*dest quede incluido explÃ­citamente en la expresiÃ³n de restricciÃ³n.
        candidate = evaluate_candidate(
            region,
            e_bar="#3",
            g_bar="#3",
            g_count=2,
            spacing_mm=160,
            long_bar="#4",
            long_count=2,
        )
        self.assertEqual(candidate.status, "fail")
        self.assertEqual(candidate.failure_mode, "region_detail_fail")
        self.assertIn("24dest", candidate.message)

    def test_transverse_feasibility_is_independent_from_longitudinal(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=4,
            width_mm=None,
            height_mm=None,
            cover_side_mm=None,
            cover_top_mm=None,
            cover_bottom_mm=None,
            source_control="gravity",
            governing_station=0.0,
            v_req=1492.33,
            t_req=1115.53,
            l_req=9999.0,
            station_count=3,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#4",
            g_bar="#3",
            g_count=2,
            spacing_mm=100,
            long_bar="#4",
            long_count=2,
        )
        self.assertEqual(candidate.status, "ok")
        self.assertEqual(candidate.failure_mode, "ok")
        self.assertGreaterEqual(candidate.at_over_s, region.t_req)
        self.assertGreaterEqual(candidate.av_over_s, region.v_req)

    def test_des_confined_spacing_rule_reports_6db_as_controlling_limit(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R1",
            region_type="C",
            beam_detailing="DES",
            d_mm=600.0,
            db_bar="#3",
            min_branches=None,
            width_mm=300.0,
            height_mm=600.0,
            cover_side_mm=40.0,
            cover_top_mm=40.0,
            cover_bottom_mm=40.0,
            source_control="gravity",
            governing_station=0.0,
            v_req=10.0,
            t_req=1.0,
            l_req=10.0,
            station_count=3,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#5",
            g_bar="#5",
            g_count=0,
            spacing_mm=70,
            long_bar="#4",
            long_count=2,
        )
        self.assertEqual(candidate.status, "fail")
        self.assertEqual(candidate.failure_mode, "region_detail_fail")
        self.assertEqual(candidate.controlling_limit, "6db")
        self.assertIn("DES region C spacing limit failed", candidate.message)

    def test_dmo_non_confined_spacing_rule_uses_ph_over_8_when_torsion_exists(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R2",
            region_type="NC",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=None,
            width_mm=300.0,
            height_mm=600.0,
            cover_side_mm=40.0,
            cover_top_mm=40.0,
            cover_bottom_mm=40.0,
            source_control="gravity",
            governing_station=0.0,
            v_req=10.0,
            t_req=1.0,
            l_req=10.0,
            station_count=3,
            fc_mpa=28.0,
            fy_mpa=420.0,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#5",
            g_bar="#3",
            g_count=0,
            spacing_mm=190,
            long_bar="#4",
            long_count=2,
        )
        self.assertEqual(candidate.status, "fail")
        self.assertEqual(candidate.failure_mode, "region_detail_fail")
        self.assertIn("Ph/8", candidate.message)

    def test_dmo_non_confined_spacing_rule_uses_d_over_4_for_high_shear(self) -> None:
        region = RegionDemand(
            beam_id="B1",
            span_id="S1",
            region_id="R2",
            region_type="NC",
            beam_detailing="DMO",
            d_mm=600.0,
            db_bar="#5",
            min_branches=None,
            width_mm=300.0,
            height_mm=600.0,
            cover_side_mm=40.0,
            cover_top_mm=40.0,
            cover_bottom_mm=40.0,
            source_control="gravity",
            governing_station=0.0,
            v_req=7000.0,
            t_req=0.0,
            l_req=10.0,
            station_count=3,
            fc_mpa=28.0,
            fy_mpa=420.0,
        )
        candidate = evaluate_candidate(
            region,
            e_bar="#5",
            g_bar="#5",
            g_count=4,
            spacing_mm=160,
            long_bar="#4",
            long_count=2,
        )
        self.assertEqual(candidate.status, "fail")
        self.assertEqual(candidate.failure_mode, "region_detail_fail")
        self.assertIn("d/4", candidate.message)

    def test_dmo_confined_region_allows_user_defined_min_branches_below_four(self) -> None:
        payload = {
            "case_name": "dmo_bad_min_branches",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DMO",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                    "d_mm": 600.0,
                                    "db_bar": "#5",
                                    "min_branches": 3,
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0, 1, 2, 3],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        cfg = validate_case_payload(payload)
        self.assertEqual(cfg.beams[0].spans[0].regions[0].min_branches, 3)

    def test_case_payload_allows_min_branches_equal_one(self) -> None:
        payload = {
            "case_name": "dmo_min_branches_one",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DMO",
                    "fc_mpa": 28.0,
                    "fy_mpa": 420.0,
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                    "d_mm": 600.0,
                                    "db_bar": "#5",
                                    "min_branches": 1,
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0, 1, 2, 3],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        cfg = validate_case_payload(payload)
        self.assertEqual(cfg.beams[0].spans[0].regions[0].min_branches, 1)

    def test_dmo_confined_region_min_branches_must_be_covered_by_g_counts(self) -> None:
        payload = {
            "case_name": "dmo_missing_g_count_capacity",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DMO",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                    "d_mm": 600.0,
                                    "db_bar": "#5",
                                    "min_branches": 6,
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0, 1, 2, 3],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        with self.assertRaises(DomainValidationError):
            validate_case_payload(payload)

    def test_enabled_optimization_requires_region_geometry(self) -> None:
        payload = {
            "case_name": "missing_geometry",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {"id": "R1", "from": 0.0, "to": 1.0, "type": "NC"},
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": True,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0, 1, 2, 3],
                    "stirrup_spacing_mm": [100],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 10,
                    "generations": 2,
                    "crossover_rate": 0.8,
                    "mutation_rate": 0.1,
                    "elite_count": 2,
                },
            },
        }
        with self.assertRaises(DomainValidationError):
            validate_case_payload(payload)

    def test_domain_validation_rejects_invalid_optimization_ranges(self) -> None:
        payload = {
            "case_name": "bad_optimization_ranges",
            "inputs": {
                "seismic_excel": "a.xlsx",
                "gravity_excel": "b.xlsx",
                "sheet_name": "Conc Bm Sum - ACI 318-08",
            },
            "units": {"rebar_per_length": "mm2/m"},
            "beams": [
                {
                    "beam_id": "B1",
                    "detailing": "DES",
                    "spans": [
                        {
                            "id": "S1",
                            "seismic": "1",
                            "gravity": "1",
                            "regions": [
                                {
                                    "id": "R1",
                                    "from": 0.0,
                                    "to": 1.0,
                                    "type": "C",
                                    "d_mm": 600.0,
                                    "db_bar": "#5",
                                    "width_mm": 300.0,
                                    "height_mm": 600.0,
                                },
                            ],
                        }
                    ],
                }
            ],
            "optimization": {
                "enabled": False,
                "objective": "min_weight",
                "variables": {
                    "E_bars": ["#3"],
                    "G_bars": ["#3"],
                    "G_counts": [0],
                    "stirrup_spacing_mm": [0],
                    "longitudinal_bars": ["#4"],
                    "longitudinal_bar_counts": [2],
                },
                "genetic_algorithm": {
                    "population_size": 3,
                    "generations": 0,
                    "crossover_rate": 1.2,
                    "mutation_rate": -0.1,
                    "elite_count": 3,
                },
            },
        }
        with self.assertRaises(DomainValidationError) as ctx:
            validate_case_payload(payload)
        fields = {issue.field for issue in ctx.exception.issues}
        self.assertIn("optimization.variables.stirrup_spacing_mm", fields)
        self.assertIn("optimization.genetic_algorithm.population_size", fields)

    def test_cli_reports_domain_validation_errors(self) -> None:
        repo = Path(__file__).resolve().parents[1]
        run_root = _new_runtime_dir("cli_invalid")
        try:
            bad_case = {
                "case_name": "cli_bad_case",
                "inputs": {
                    "seismic_excel": str((repo / "examples" / "case_0001" / "sismo.xlsx").resolve()),
                    "gravity_excel": str((repo / "examples" / "case_0001" / "gravedad.xlsx").resolve()),
                    "sheet_name": "Conc Bm Sum - ACI 318-08",
                },
                "units": {"rebar_per_length": "mm2/m"},
                "beams": [
                    {
                        "beam_id": "B1",
                        "detailing": "DMO",
                        "cover_side_mm": 40.0,
                        "cover_top_mm": 40.0,
                        "cover_bottom_mm": 40.0,
                        "fc_mpa": 28.0,
                        "fy_mpa": 420.0,
                        "spans": [
                            {
                                "id": "S1",
                                "seismic": "190",
                                "gravity": "190",
                                "regions": [
                                    {
                                        "id": "R1",
                                        "from": 0.0,
                                        "to": 1.0,
                                        "type": "C",
                                        "d_mm": 600.0,
                                        "db_bar": "#6",
                                        "min_branches": 3,
                                        "width_mm": 300.0,
                                        "height_mm": 600.0,
                                    }
                                ],
                            }
                        ],
                    }
                ],
                "optimization": {
                    "enabled": False,
                    "objective": "min_weight",
                    "variables": {
                        "E_bars": ["#3"],
                        "G_bars": ["#3"],
                        "G_counts": [0, 1, 2, 3],
                        "stirrup_spacing_mm": [100],
                        "longitudinal_bars": ["#4"],
                        "longitudinal_bar_counts": [2],
                    },
                    "genetic_algorithm": {
                        "population_size": 3,
                        "generations": 2,
                        "crossover_rate": 0.8,
                        "mutation_rate": 0.1,
                        "elite_count": 2,
                    },
                },
            }
            case_path = run_root / "bad_case.json"
            case_path.write_text(json.dumps(bad_case), encoding="utf-8")

            process = subprocess.run(
                [sys.executable, "-m", "rc_shear_torsion.run", str(case_path), "--out", str(run_root / "out")],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(process.returncode, 2)
            self.assertIn("Domain validation failed", process.stderr)
            self.assertIn("population_size", process.stderr)
        finally:
            shutil.rmtree(run_root, ignore_errors=True)

    def test_resolve_path_prefers_case_directory_over_cwd(self) -> None:
        root = _new_runtime_dir("resolve_prefers")
        try:
            case_dir = root / "case"
            cwd_dir = root / "cwd"
            case_inputs = case_dir / "inputs"
            cwd_inputs = cwd_dir / "inputs"
            case_inputs.mkdir(parents=True, exist_ok=True)
            cwd_inputs.mkdir(parents=True, exist_ok=True)
            case_file = case_inputs / "data.xlsx"
            cwd_file = cwd_inputs / "data.xlsx"
            case_file.write_text("case", encoding="utf-8")
            cwd_file.write_text("cwd", encoding="utf-8")

            previous_cwd = Path.cwd()
            os.chdir(cwd_dir)
            try:
                resolved = resolve_path("inputs/data.xlsx", case_dir)
            finally:
                os.chdir(previous_cwd)

            self.assertEqual(resolved, case_file.resolve())
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_resolve_path_falls_back_to_cwd(self) -> None:
        root = _new_runtime_dir("resolve_fallback")
        try:
            case_dir = root / "case"
            cwd_dir = root / "cwd"
            case_dir.mkdir(parents=True, exist_ok=True)
            cwd_inputs = cwd_dir / "inputs"
            cwd_inputs.mkdir(parents=True, exist_ok=True)
            cwd_file = cwd_inputs / "data.xlsx"
            cwd_file.write_text("cwd", encoding="utf-8")

            previous_cwd = Path.cwd()
            os.chdir(cwd_dir)
            try:
                resolved = resolve_path("inputs/data.xlsx", case_dir)
            finally:
                os.chdir(previous_cwd)

            self.assertEqual(resolved, cwd_file.resolve())
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_resolve_path_missing_is_anchored_to_case_directory(self) -> None:
        root = _new_runtime_dir("resolve_missing")
        try:
            case_dir = root / "case"
            cwd_dir = root / "cwd"
            case_dir.mkdir(parents=True, exist_ok=True)
            cwd_dir.mkdir(parents=True, exist_ok=True)

            previous_cwd = Path.cwd()
            os.chdir(cwd_dir)
            try:
                resolved = resolve_path("inputs/missing.xlsx", case_dir)
            finally:
                os.chdir(previous_cwd)

            self.assertEqual(resolved, (case_dir / "inputs" / "missing.xlsx").resolve())
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_end_to_end_case_runs(self) -> None:
        repo = Path(__file__).resolve().parents[1]
        case_json = repo / "examples" / "case_0001" / "case.json"
        run_root = _new_runtime_dir("end_to_end")
        try:
            out_root = run_root / "results"
            out_dir = run_case(case_json, out_root)
            self.assertTrue((out_dir / "design_results.xlsx").exists())
            self.assertTrue((out_dir / "summary.xlsx").exists())
            self.assertTrue((out_dir / "optimized_results.xlsx").exists())
            self.assertTrue((out_dir / "reinforcement_schedule.xlsx").exists())
            self.assertTrue((out_dir / "run_log.txt").exists())

            wb = load_workbook(out_dir / "reinforcement_schedule.xlsx", data_only=True)
            self.assertIn("por_region", wb.sheetnames)
            por_region = wb["por_region"]
            # 3 regiones en case_0001 con al menos 5 opciones por region.
            self.assertGreaterEqual(por_region.max_row - 1, 15)
            schedule_headers = [cell.value for cell in por_region[1]]
            self.assertEqual(
                schedule_headers,
                [
                    "viga_id",
                    "vano_id",
                    "region_id",
                    "opcion",
                    "longitud_region_mm",
                    "estado",
                    "arreglo_transversal",
                    "limite_controlante",
                    "cantidad_estribos_region",
                    "arreglo_longitudinal",
                    "peso_unitario_estribo_kg",
                    "peso_transversal_region_kg",
                    "peso_longitudinal_region_kg",
                    "peso_total_region_kg",
                ],
            )

            por_vano = wb["por_vano"]
            por_vano_headers = [cell.value for cell in por_vano[1]]
            self.assertEqual(
                por_vano_headers,
                [
                    "viga_id",
                    "vano_id",
                    "regiones_totales",
                    "regiones_cumplen",
                    "regiones_fallan",
                    "peso_transversal_total_kg",
                    "peso_longitudinal_total_kg",
                    "peso_total_kg",
                ],
            )

            optimized_wb = load_workbook(out_dir / "optimized_results.xlsx", data_only=True)
            optimized = optimized_wb.active
            optimized_headers = [cell.value for cell in optimized[1]]
            self.assertEqual(
                optimized_headers,
                [
                    "beam_id",
                    "span_id",
                    "region_id",
                    "method",
                    "status",
                    "failure_mode",
                    "objective",
                    "source_control",
                    "governing_station",
                    "E_bar",
                    "G_bar",
                    "G_count",
                    "spacing_mm",
                    "controlling_limit",
                    "long_bar",
                    "long_count",
                    "long_provided_mm2",
                    "VRebar_req",
                    "TTrnRebar_req",
                    "TLngRebar_req",
                    "VRebar_req_units",
                    "TTrnRebar_req_units",
                    "TLngRebar_req_units",
                    "check_torsion",
                    "check_shear",
                    "check_longitudinal",
                    "check_detailing",
                    "transverse_weight_kg_per_m",
                    "longitudinal_weight_kg_per_m",
                    "total_weight_kg_per_m",
                    "evaluated_candidates",
                    "feasible_candidates",
                ],
            )

            design_wb = load_workbook(out_dir / "design_results.xlsx", data_only=True)
            design = design_wb.active
            design_headers = [cell.value for cell in design[1]]
            self.assertEqual(
                design_headers,
                [
                    "beam_id",
                    "span_id",
                    "region_id",
                    "region_type",
                    "source_control",
                    "governing_station",
                    "VRebar_req",
                    "VRebar_req_units",
                    "TTrnRebar_req",
                    "TTrnRebar_req_units",
                    "TLngRebar_req",
                    "TLngRebar_req_units",
                    "E_bar",
                    "G_bar",
                    "G_count",
                    "spacing_mm",
                    "controlling_limit",
                    "Av1",
                    "Av2",
                    "Av_total",
                    "At",
                    "At_over_s",
                    "Av_over_s",
                    "long_bar",
                    "long_count",
                    "long_provided_mm2",
                    "check_torsion",
                    "check_shear",
                    "check_longitudinal",
                    "check_detailing",
                    "failure_mode",
                    "status",
                    "message",
                ],
            )

            summary_wb = load_workbook(out_dir / "summary.xlsx", data_only=True)
            span_summary_headers = [cell.value for cell in summary_wb["span_summary"][1]]
            self.assertEqual(
                span_summary_headers,
                ["beam_id", "span_id", "total_regions", "ok_regions", "fail_regions", "status", "message"],
            )
            beam_summary_headers = [cell.value for cell in summary_wb["beam_summary"][1]]
            self.assertEqual(
                beam_summary_headers,
                ["beam_id", "total_spans", "ok_spans", "fail_spans", "status"],
            )
        finally:
            shutil.rmtree(run_root, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()

